import tempfile
import json
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

CHOICES = {"production_scheme": "image_full_rebuild", "reconstruction_mode": "preserve_complex_images"}


def fixture_render_receipt(project):
    """Bind synthetic fixture images; production receipts are written by the renderer."""
    from scripts.artifact_receipt import fingerprint
    folder = project / 'reports/rendered_png'
    (folder/'render_receipt.json').write_text(json.dumps({
        'status':'pass', 'pptx_identity':fingerprint(project/'pptx/output.pptx'),
        'render_identities':[fingerprint(p) for p in sorted(folder.glob('slide_*.png'))],
    }), encoding='utf-8')


def complete_project(root):
    from scripts.image_reconstruction_pipeline import init_project
    source = root / "source.png"
    Image.new("RGB", (320, 180), "white").save(source)
    project = root / "project"
    init_project(project, [source], **CHOICES)
    path = project / "analysis/_analysis.json"
    inventory = json.loads(path.read_text(encoding="utf-8"))
    inventory["slides"][0]["completeness_check"] = {"performed": True, "layer_a_count": 0}
    inventory["slides"][0]["elements"] = [{"element_id": "title", "layer": "C", "implementation": "native_text", "text": "Title", "bbox_percent": {"x": 10, "y": 10, "w": 30, "h": 10}}]
    inventory["slides"][0]["source_text_lines"] = ["Title"]
    path.write_text(json.dumps(inventory), encoding="utf-8")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # The reconstruction contract maps source percentages onto the 1280x720
    # canvas.  Keep the fixture's native frame on that same geometry.
    slide.shapes.add_textbox(Inches(1), Inches(0.5625), Inches(3), Inches(0.5625)).text = "Title"
    prs.save(project / "pptx/output.pptx")
    rendered = project / "reports/rendered_png"
    rendered.mkdir()
    Image.new("RGB", (320, 180), "black").save(rendered / "slide_001.png")
    fixture_render_receipt(project)
    return project


class ImageReconstructionPipelineTests(unittest.TestCase):
    def test_reinitialization_never_replaces_sources(self):
        from scripts.image_reconstruction_pipeline import init_project
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            a, b = root / "a.png", root / "b.png"
            Image.new("RGB", (32, 18), "red").save(a)
            Image.new("RGB", (32, 18), "blue").save(b)
            project = root / "project"
            init_project(project, [a], **CHOICES)
            original = (project / "sources/slide_001.png").read_bytes()
            analysis = (project / "analysis/_analysis.json").read_bytes()
            for overwrite in (False, True):
                with self.assertRaises(FileExistsError):
                    init_project(project, [b], overwrite_analysis=overwrite, **CHOICES)
                self.assertEqual((project / "sources/slide_001.png").read_bytes(), original)
                self.assertEqual((project / "analysis/_analysis.json").read_bytes(), analysis)
            init_project(project, [a], overwrite_analysis=True, **CHOICES)

    def test_invalid_later_source_does_not_copy_earlier_source(self):
        from scripts.image_reconstruction_pipeline import init_project
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.png"
            Image.new("RGB", (32, 18)).save(source)
            with self.assertRaises(FileNotFoundError):
                init_project(root / "project", [source, root / "missing.png"], **CHOICES)
            self.assertFalse((root / "project/sources/slide_001.png").exists())

    def test_init_project_creates_canonical_scaffold(self):
        from scripts.image_reconstruction_pipeline import init_project

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.png"
            Image.new("RGB", (320, 180), "white").save(source)
            project = root / "image_project"

            report = init_project(project, [source], **CHOICES)

            self.assertEqual(report["status"], "initialized")
            self.assertTrue((project / "sources" / "slide_001.png").exists())
            self.assertTrue((project / "analysis" / "_analysis.json").exists())
            self.assertTrue((project / "pages" / "page_001" / "assets" / "split").is_dir())
            self.assertTrue((project / "pptx").is_dir())
            self.assertTrue((project / "reports").is_dir())
            inventory = json.loads((project / "analysis" / "_analysis.json").read_text(encoding="utf-8"))
            self.assertEqual(inventory["alignment_contract"]["schema_version"], "easyslides.alignment_contract.v1")

    def test_practical_mode_treats_source_diff_as_advisory(self):
        from scripts.image_reconstruction_pipeline import qa_project

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            project = root / "image_project"
            rendered = project / "reports" / "rendered_png"
            rendered.mkdir(parents=True)
            source = root / "source.png"
            Image.new("RGB", (320, 180), "white").save(source)
            Image.new("RGB", (320, 180), "black").save(rendered / "slide_001.png")

            report = qa_project(
                project,
                rendered_dir=rendered,
                source_images=[source],
                inventory=project / "analysis" / "missing.json",
                mode="faithful-practical",
            )

            self.assertEqual(report["status"], "fail")
            diff_gate = next(gate for gate in report["gates"] if gate["name"] == "source_render_diff")
            self.assertEqual(diff_gate["status"], "fail")
            self.assertTrue(diff_gate["advisory"])
            self.assertEqual(diff_gate["blocking_count"], 0)
            self.assertGreater(diff_gate["raw_blocking_count"], 0)

    def test_complete_inputs_practical_pass_and_strict_fail(self):
        from scripts.image_reconstruction_pipeline import qa_project
        with tempfile.TemporaryDirectory() as tmp:
            project = complete_project(Path(tmp))
            practical = qa_project(project, mode='faithful-practical')
            self.assertEqual(practical["status"], "pass", practical)
            self.assertFalse(practical['delivery_ready'])
            self.assertEqual(qa_project(project, mode="pixel-strict")["status"], "fail")

    def test_missing_pptx_inventory_or_render_always_blocks(self):
        from scripts.image_reconstruction_pipeline import qa_project
        for path in ("pptx/output.pptx", "analysis/_analysis.json", "reports/rendered_png/slide_001.png"):
            with tempfile.TemporaryDirectory() as tmp:
                project = complete_project(Path(tmp))
                (project / path).unlink()
                self.assertEqual(qa_project(project, mode="pixel-strict")["status"], "fail", path)

    def test_source_reviewed_lines_required_for_text_acceptance(self):
        from scripts.image_reconstruction_pipeline import qa_project
        for value in (None, [], "Title", [""], [3], ["Title\nSubtitle"], ["Other"]):
            with self.subTest(value=value), tempfile.TemporaryDirectory() as tmp:
                project = complete_project(Path(tmp))
                path = project / "analysis/_analysis.json"
                inventory = json.loads(path.read_text(encoding="utf-8"))
                if value is None:
                    del inventory["slides"][0]["source_text_lines"]
                else:
                    inventory["slides"][0]["source_text_lines"] = value
                path.write_text(json.dumps(inventory), encoding="utf-8")
                self.assertEqual(qa_project(project)["status"], "fail")
                report = json.loads((project / "reports/image_reconstruction_pptx_report.json").read_text(encoding="utf-8"))
                self.assertTrue(any(issue["code"] == "PPTX-SOURCE-TEXT-LINES" for issue in report["issues"]))

    def test_extra_render_and_missing_choices_block(self):
        from scripts.image_reconstruction_pipeline import qa_project, init_project
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            project = complete_project(root)
            Image.new("RGB", (320, 180)).save(project / "reports/rendered_png/slide_002.png")
            self.assertEqual(qa_project(project)["status"], "fail")
            with self.assertRaises(ValueError):
                init_project(root / "unconfirmed", [root / "source.png"])
            self.assertFalse((root / "unconfirmed").exists())

    def test_partial_rebuild_keeps_unselected_pages_and_detects_changed_background(self):
        from scripts.image_reconstruction_pipeline import qa_project
        with tempfile.TemporaryDirectory() as tmp:
            project = complete_project(Path(tmp))
            path = project / "analysis/_analysis.json"
            inventory = json.loads(path.read_text())
            inventory["production_scheme"] = "image_partial_rebuild"
            inventory["slides"][0]["reconstruction_regions"] = [{"x": 10, "y": 10, "w": 30, "h": 10}]
            inventory["slides"].append({"slide_id": "s02", "source_image": "sources/slide_002.png", "elements": [], "reconstruction_regions": [], "completeness_check": {"performed": True, "layer_a_count": 0}})
            path.write_text(json.dumps(inventory), encoding="utf-8")
            image = Image.new("RGB", (320, 180), "white")
            image.save(project / "sources/slide_002.png")
            for name in ("slide_001.png", "slide_002.png"):
                image.save(project / "reports/rendered_png" / name)
            prs = Presentation(project / "pptx/output.pptx")
            prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(str(project / "sources/slide_002.png"), 0, 0, width=prs.slide_width, height=prs.slide_height).name = "source_background"
            prs.save(project / "pptx/output.pptx")
            from scripts.artifact_receipt import fingerprint
            run_path = project/'image_reconstruction_run.json'
            run = json.loads(run_path.read_text(encoding='utf-8'))
            run['decisions']['production_scheme'] = 'image_partial_rebuild'
            run['source_identities'] = [fingerprint(project/'sources'/name) for name in ('slide_001.png','slide_002.png')]
            run_path.write_text(json.dumps(run), encoding='utf-8')
            fixture_render_receipt(project)
            self.assertEqual(qa_project(project)["status"], "pass")
            Image.new("RGB", (320, 180), "black").save(project / "reports/rendered_png/slide_002.png")
            fixture_render_receipt(project)
            report = qa_project(project)
            self.assertEqual(report["status"], "fail")
            self.assertTrue(any(i["code"] == "QA-UNSELECTED-CHANGED" for i in json.loads((project / "reports/required_inputs_report.json").read_text())["issues"]))

    def test_pixel_strict_mode_blocks_on_source_diff(self):
        from scripts.image_reconstruction_pipeline import qa_project

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            project = root / "image_project"
            rendered = project / "reports" / "rendered_png"
            rendered.mkdir(parents=True)
            source = root / "source.png"
            Image.new("RGB", (320, 180), "white").save(source)
            Image.new("RGB", (320, 180), "black").save(rendered / "slide_001.png")

            report = qa_project(
                project,
                rendered_dir=rendered,
                source_images=[source],
                inventory=project / "analysis" / "missing.json",
                mode="pixel-strict",
            )

            self.assertEqual(report["status"], "fail")
            diff_gate = next(gate for gate in report["gates"] if gate["name"] == "source_render_diff")
            self.assertFalse(diff_gate["advisory"])
            self.assertGreater(diff_gate["blocking_count"], 0)

    def test_authored_page_svg_alignment_gate_blocks_parent_transform(self):
        from scripts.image_reconstruction_pipeline import qa_project

        with tempfile.TemporaryDirectory() as tmp:
            project = complete_project(Path(tmp))
            page_svg = project / "pages/page_001/slide_001.svg"
            page_svg.write_text(
                """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720">
  <g transform="translate(40 0)">
    <text data-pptx-textbox="true" data-pptx-box-x="128" data-pptx-box-y="72" data-pptx-box-w="384" data-pptx-box-h="72" x="320" y="108" text-anchor="middle" data-pptx-valign="middle">Title</text>
  </g>
</svg>""",
                encoding="utf-8",
            )

            report = qa_project(project)

        gate = next(gate for gate in report["gates"] if gate["name"] == "alignment_contract_svg")
        self.assertEqual(gate["status"], "fail")
        self.assertGreater(gate["blocking_count"], 0)


if __name__ == "__main__":
    unittest.main()
