import json
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]


class VisualIntelligenceTests(unittest.TestCase):
    def test_all_official_templates_resolve_semantic_theme_tokens(self):
        from scripts.theme_tokens import REQUIRED_ROLES, resolve_template_tokens, validate_tokens

        policy = json.loads((ROOT / "templates" / "template_policy.json").read_text(encoding="utf-8"))
        for template_id in policy["official_template_ids"]:
            resolved = resolve_template_tokens(ROOT / "templates" / "layouts" / template_id)
            self.assertEqual(validate_tokens(resolved["tokens"])["status"], "pass", template_id)
            self.assertEqual(set(REQUIRED_ROLES), set(resolved["tokens"]))

    def test_visual_planner_preserves_evidence_and_plans_hero_pages(self):
        from scripts.visual_asset_planner import plan_visual_assets

        plan = plan_visual_assets(
            {
                "deck_id": "demo",
                "slides": [
                    {"page": "P01", "role": "cover", "action_title": "生态系统恢复"},
                    {"page": "P02", "role": "chapter", "action_title": "研究方法"},
                    {"page": "P03", "role": "content", "content_shape": "figure_evidence", "evidence_sources": [{"source_id": "fig-1"}]},
                ],
            },
            policy="auto_decorative",
            imagegen_available=True,
        )
        self.assertEqual(plan["planned_count"], 2)
        self.assertEqual(
            [row["decision"] for row in plan["decisions"]],
            ["hero_background", "hero_background", "source_only"],
        )
        self.assertTrue(all("No visible words" in item["prompt"] for item in plan["items"]))

        themed = plan_visual_assets(
            {"slides": [{"page": "P01", "role": "cover", "title": "Theme"}]},
            template_dir=ROOT / "templates" / "layouts" / "academic_general",
            palette_id="academic_green",
            imagegen_available=True,
        )
        self.assertEqual(themed["theme_tokens"]["primary"], "#146B4A")

    def test_transition_pages_are_eligible_for_hero_backgrounds(self):
        from scripts.visual_asset_planner import plan_visual_assets

        plan = plan_visual_assets(
            {"slides": [{"page": "P02", "role": "transition", "title": "Methods"}]},
            policy="auto_decorative",
            imagegen_available=True,
        )
        self.assertEqual(plan["planned_count"], 1)
        self.assertEqual(plan["items"][0]["slide_role"], "transition")

    def test_ai_rich_hero_policy_can_plan_background_and_local_decorative_art(self):
        from scripts.visual_asset_planner import plan_visual_assets

        plan = plan_visual_assets(
            {"slides": [{"page": "C01", "role": "cover", "title": "A quiet cover"}]},
            template_dir=ROOT / "templates" / "layouts" / "literature_minimal",
            policy="ai_rich",
            imagegen_available=True,
        )
        self.assertEqual(plan["planned_count"], 2)
        self.assertEqual(plan["decisions"][0]["decision"], "hero_background+local_illustration")
        self.assertEqual(
            {item["asset_role"] for item in plan["items"]},
            {"background", "illustration"},
        )
        self.assertTrue(all(item["source_policy"] == "generated_decorative_only" for item in plan["items"]))

    def test_template_shell_readability_policy_keeps_opaque_ending_band_light(self):
        from scripts.visual_asset_planner import plan_visual_assets

        plan = plan_visual_assets(
            {"slides": [{"page": "E01", "role": "ending", "title": "Closing"}]},
            template_dir=ROOT / "templates" / "layouts" / "literature_minimal",
            policy="ai_rich",
            imagegen_available=True,
        )
        background = next(item for item in plan["items"] if item["asset_role"] == "background")
        self.assertEqual(background["background"]["readability_mode"], "template_shell")
        self.assertEqual(background["background"]["text_tone"], "light")
        self.assertEqual(background["background"]["text_color"], "#FFFFFF")

        illustration = next(item for item in plan["items"] if item["asset_role"] == "illustration")
        self.assertEqual(illustration["placement"]["layer"], "shell_overlay")

    def test_user_theme_tokens_override_template_tokens_and_recolor_svg(self):
        from scripts.slide_compiler import _slide_theme
        from scripts.template_compiler import compile_template

        template_ir = compile_template("literature_minimal", palette_id="literature_blue")["template_ir"]
        theme = _slide_theme(
            {"theme_tokens": {"primary": "#7A1FA2", "accent": "#7A1FA2", "text_on_light": "#102030"}},
            template_ir,
        )
        self.assertEqual(theme["tokens"]["primary"], "#7A1FA2")
        self.assertEqual(theme["tokens"]["text_on_light"], "#102030")
        self.assertIn("#0D5DBE", theme["replacements"])
        self.assertEqual(theme["replacements"]["#0D5DBE"], "#7A1FA2")
        from scripts.slide_compiler import SlideCompileError
        with self.assertRaises(SlideCompileError):
            _slide_theme({"theme_tokens": {"text_on_dark": "#FF0000"}}, template_ir)

    def test_visual_planner_can_write_manifest_and_bind_deck_plan(self):
        from scripts.visual_asset_planner import write_visual_plan

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            deck_path = root / "deck_plan.json"
            deck_path.write_text(
                json.dumps({"slides": [{"page": "P01", "role": "cover", "title": "Demo"}]}),
                encoding="utf-8",
            )
            result = write_visual_plan(deck_path, root / "images", imagegen_available=False, write_deck_plan=True)
            self.assertEqual(result["planned_count"], 1)
            updated = json.loads(deck_path.read_text(encoding="utf-8"))
            self.assertIn("image_manifest", updated)
            self.assertEqual(updated["slides"][0]["image_resource_id"], "p01_background")
            self.assertTrue((root / "images" / "image_prompts.json").is_file())

    def test_readability_chooses_light_for_dark_and_dark_for_light(self):
        from scripts.readability_gate import choose_text_tone

        dark = choose_text_tone(0.03)
        light = choose_text_tone(0.92)
        self.assertEqual(dark["tone"], "light")
        self.assertEqual(light["tone"], "dark")
        self.assertTrue(dark["passes"])
        self.assertTrue(light["passes"])

    def test_readability_detects_mixed_background_and_requests_scrim(self):
        from PIL import Image
        from scripts.readability_gate import analyze_region

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "mixed.png"
            image = Image.new("RGB", (200, 100), "white")
            for x in range(100):
                for y in range(100):
                    image.putpixel((x, y), (0, 0, 0))
            image.save(path)
            report = analyze_region(image, {"box": {"x": 0, "y": 0, "width": 200, "height": 100}})
            self.assertTrue(report["mixed_background"])
            self.assertTrue(report["needs_scrim"])

    def test_build_facade_preserves_explicit_source_images_and_writes_a_report(self):
        from scripts.build_pipeline import build_parser, run_build
        from scripts.image_acquisition import scaffold_hero_manifest

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest_path = root / "images" / "image_prompts.json"
            scaffold_hero_manifest(manifest_path, topic="source image preservation")
            deck_path = root / "deck_plan.json"
            deck_path.write_text(
                json.dumps(
                    {
                        "template_id": "academic_general",
                        "production_scheme": "direct_editable",
                        "image_manifest": "images/image_prompts.json",
                        "slides": [
                            {
                                "page": "P01",
                                "role": "cover",
                                "image_resource_id": "cover_background",
                                "shell_payload": {"LOGO": "EasySlides", "TITLE": "Build", "SUBTITLE": "Test", "AUTHOR": "A", "ADVISOR": "B", "INSTITUTION": "C", "DATE": "2026"},
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            args = build_parser().parse_args([str(deck_path), "--out-dir", str(root / "build"), "--no-svg"])
            report = run_build(args)
            self.assertEqual(report["status"], "pass")
            build_plan = json.loads((root / "build" / "deck_plan.build.json").read_text(encoding="utf-8"))
            self.assertEqual(build_plan["slides"][0]["image_resource_id"], "cover_background")
            merged = json.loads((root / "build" / "visual_assets" / "image_prompts.json").read_text(encoding="utf-8"))
            self.assertIn("cover_background", {item["id"] for item in merged["items"]})
            self.assertEqual(report["asset_readiness"], "pending_or_manual")

    def test_build_requires_scheme_before_creating_output(self):
        from scripts.build_pipeline import build_parser, run_build
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            plan = root / "plan.json"
            plan.write_text(json.dumps({"slides": []}), encoding="utf-8")
            with self.assertRaises(ValueError):
                run_build(build_parser().parse_args([str(plan), "--out-dir", str(root / "build")]))
            self.assertFalse((root / "build").exists())

    def test_build_no_images_and_bound_images_export_native_text_with_palette(self):
        from PIL import Image
        from pptx import Presentation
        from scripts.build_pipeline import build_parser, run_build
        from scripts.image_acquisition import scaffold_hero_manifest
        from xml.etree import ElementTree as ET

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            plan = {"template_id": "academic_general", "production_scheme": "direct_editable",
                    "slides": [{"page": "C01", "role": "cover", "shell_payload": {
                        "LOGO": "LAB", "TITLE": "Native title", "SUBTITLE": "Subtitle",
                        "AUTHOR": "A", "ADVISOR": "B", "INSTITUTION": "C", "DATE": "2026"}}]}
            path = root / "deck.json"
            path.write_text(json.dumps(plan), encoding="utf-8")
            report = run_build(build_parser().parse_args([str(path), "--out-dir", str(root / "plain"),
                                                         "--no-visual-plan", "--palette", "academic_green"]))
            self.assertEqual(report["status"], "pass")
            self.assertEqual(report["acquisition"]["status"], "skipped")
            svg = ET.parse(report["svg_render"]["svg_files"][0]).getroot()
            self.assertTrue(any(node.get("fill") == "#146B4A" for node in svg.iter()))
            manifest = root / "images" / "image_prompts.json"
            scaffold_hero_manifest(manifest, topic="Background")
            Image.new("RGB", (1280, 720), "white").save(manifest.parent / "cover_bg.png")
            plan["image_manifest"] = "images/image_prompts.json"
            path.write_text(json.dumps(plan), encoding="utf-8")
            output = root / "deck.pptx"
            report = run_build(build_parser().parse_args([str(path), "--out-dir", str(root / "pictured"),
                                                         "--no-visual-plan", "--pptx-out", str(output)]))
            self.assertEqual(report["status"], "pass")
            shapes = Presentation(output).slides[0].shapes
            self.assertTrue(Presentation(output).slides[0]._element.xpath(".//p:pic"))
            self.assertTrue(any(shape.has_text_frame and shape.text == "Native title" for shape in shapes))


if __name__ == "__main__":
    unittest.main()
