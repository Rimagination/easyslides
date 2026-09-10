"""Small offline regressions for the annotation-assisted figure experiment."""
from copy import deepcopy
from pathlib import Path
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET

import cv2
import numpy as np
from pptx import Presentation
from pptx.oxml import parse_xml

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from rebuild_bitmap_figure import (
    add_label, assemble_scene, prepare_layers, scaled_box, validate_annotations,
)
from svg_to_pptx.drawingml_converter import convert_svg_to_slide_shapes


class FigureRebuildTests(unittest.TestCase):
    def test_annotation_validation_and_scientific_text(self):
        data = {'reference_size':[100,100], 'tree_coordinate_size':[100,100],
                'trees':[], 'labels':[{'id':'nitrate', 'box':[10,10,50,40],
                                      'text':'NO3-', 'runs':[['NO',0],['3',-1],['-',1]]}]}
        validate_annotations(data, (200, 200))
        self.assertEqual(scaled_box([10,10,50,40], [100,100], [200,200]), [20,20,100,80])
        invalid = deepcopy(data)
        invalid['labels'].append(deepcopy(invalid['labels'][0]))
        with self.assertRaises(ValueError):
            validate_annotations(invalid, (200,200))
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        label = add_label(slide, data['labels'][0], data, (200,200))
        self.assertEqual(label.text, 'NO3-')
        self.assertEqual([r.get('baseline') for r in label._element.xpath('.//a:rPr')], [None,'-22000','40000'])

    def test_no_subpath_splitting_or_background_reordering(self):
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            # Reverse inner winding makes a real hole, plus a later overlay.
            (work/'background.svg').write_text(
                '<svg xmlns="http://www.w3.org/2000/svg"><path id="ring" fill="#ffffff" '
                'd="M0 0H80V80H0Z M20 20V60H60V20Z"/><rect id="later" x="30" y="30" '
                'width="5" height="5" fill="#ff0000"/></svg>')
            trees = [{'id':'tree', 'points':[[90,5],[85,20],[95,20]], 'fill':'#338833'}]
            assemble_scene(work/'background.svg', trees, work/'scene.svg', (100,100))
            group = list(ET.parse(work/'scene.svg').getroot())[0]
            self.assertEqual([p.get('id') for p in group], ['ring','later'])
            xml, media, _, _ = convert_svg_to_slide_shapes(work/'scene.svg')
            root = parse_xml(xml.encode())
            ring = root.xpath('.//p:grpSp/p:sp')[0]
            self.assertEqual(len(ring.xpath('.//a:path')), 1)
            self.assertEqual(len(ring.xpath('.//a:moveTo')), 2)
            self.assertEqual(len(ring.xpath('.//a:close')), 2)
            self.assertFalse(media)

    def test_tree_removed_before_trace_and_outside_pixels_preserved(self):
        source = np.full((100,100,3), (100,180,160), dtype=np.uint8)
        cv2.fillPoly(source, [np.array([[50,20],[35,65],[65,65]])], (60,120,60))
        data = {'reference_size':[100,100], 'tree_coordinate_size':[100,100],
                'labels':[], 'trees':[{'id':'test_tree','box':[30,15,70,70]}]}
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            trees = prepare_layers(source, data, work)
            self.assertEqual(len(trees), 1)
            mask = cv2.imread(str(work/'tree_mask.png'), 0)
            background = cv2.imread(str(work/'background.png'))
            np.testing.assert_array_equal(source[mask == 0], background[mask == 0])
            self.assertGreater(int(background[45,50,1]), 160)


if __name__ == '__main__':
    unittest.main()
