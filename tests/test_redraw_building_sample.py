"""Small geometric and native-editability checks, no tracing engine required."""
from pathlib import Path
import sys
import tempfile
import unittest

from PIL import Image
from pptx import Presentation

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from redraw_building_sample import quad, window_grid, paint


class ArchitecturalRedrawTests(unittest.TestCase):
    def setUp(self):
        self.spec={'name':'Building','origin':[10,10],'u':[80,28], 'v':[40,-14], 'height':90}
        self.image=Image.new('RGB',(160,160),'white')

    def test_regular_windows_share_facade_axes_and_pitch(self):
        group=window_grid(self.spec,{'side':'left','cols':5,'rows':4,'glass':'8799AB'},self.image)
        self.assertEqual(len(group['children']),20)
        widths=[]
        for window in group['children']:
            p=window['paths'][0]
            self.assertEqual(len(p),4)
            self.assertAlmostEqual((p[1][1]-p[0][1])/(p[1][0]-p[0][0]),28/80)
            self.assertAlmostEqual(p[0][0],p[3][0])
            self.assertAlmostEqual(p[1][0],p[2][0])
            widths.append(p[1][0]-p[0][0])
        self.assertLess(max(widths)-min(widths),1e-8)
        first_row=[w['paths'][0][0][0] for w in group['children'][:5]]
        steps=[b-a for a,b in zip(first_row,first_row[1:])]
        self.assertLess(max(steps)-min(steps),1e-8)

    def test_four_panes_are_one_flat_native_window_after_save(self):
        group=window_grid(self.spec,{'side':'left','cols':1,'rows':1,'glass':'8799AB','pane_grid':True},self.image)
        self.assertEqual(len(group['children'][0]['paths']),4)
        prs=Presentation()
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        paint(slide.shapes,[{'name':'Building','children':[group]}])
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'window.pptx'
            prs.save(path)
            loaded=Presentation(path).slides[0]
            self.assertEqual(len(loaded.shapes),1)
            native=loaded._element.xpath('.//p:sp')
            self.assertEqual(len(native),1)
            self.assertEqual(len(native[0].xpath('.//a:moveTo')),4)
            self.assertEqual(len(native[0].xpath('.//a:lnTo')),12)
            self.assertEqual(len(native[0].xpath('.//a:close')),4)
            self.assertEqual(len(native[0].xpath('./p:spPr/a:effectLst')),1)
            self.assertFalse(loaded._element.xpath('.//p:pic'))

    def test_invalid_geometry_rejected(self):
        with self.assertRaises(ValueError):
            quad([0,0],[10,4],20,.8,0,.2,1)
        with self.assertRaises(ValueError):
            window_grid(self.spec,{'side':'left','cols':0,'rows':4,'glass':'8799AB'},self.image)


if __name__=='__main__':
    unittest.main()
