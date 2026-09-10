"""Small artifact regression check for hybrid geometry, text, and movement."""
from pathlib import Path
import json
import unittest

from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
OUTPUT=ROOT/'output/bitmap_trace_test01/hybrid'


@unittest.skipUnless(OUTPUT.is_dir(), "Optional local hybrid acceptance artifacts are not installed")
class HybridTests(unittest.TestCase):
    def test_hybrid_editability_and_saved_movement(self):
        source=Presentation(OUTPUT/'test01_hybrid.pptx')
        demo=Presentation(OUTPUT/'test01_hybrid_edit_demo.pptx')
        data=json.loads((ROOT/'output/bitmap_trace_test01/editable_v2/annotations.json').read_text(encoding='utf-8'))
        report=json.loads((OUTPUT/'verification.json').read_text(encoding='utf-8'))
        original={s.name:s for s in source.slides[0].shapes}
        moved={s.name:s for s in demo.slides[0].shapes}
        self.assertIn('Warehouse',original)
        self.assertIn('Traced_background_details',original)
        self.assertEqual(original['Traced_background_details']._element.xml,
                         moved['Traced_background_details']._element.xml)
        tree=report['moved_tree']
        self.assertGreater(moved[tree].left,original[tree].left)
        self.assertEqual(moved[tree].top,original[tree].top)
        self.assertEqual(moved['Text_industry'].text,'Editable text')
        for label in data['labels']:
            self.assertEqual(original['Text_'+label['id']].text,label['text'])
        for deck in (source,demo):
            xml=deck.slides[0]._element
            self.assertFalse(xml.xpath('.//p:pic'))
            identifiers=xml.xpath('.//p:cNvPr/@id')
            self.assertEqual(len(identifiers),len(set(identifiers)))
        self.assertEqual(sum(s.startswith('Tree_region_') for s in original),report['tree_regions'])


if __name__=='__main__':
    unittest.main()
