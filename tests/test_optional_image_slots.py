import json
from pathlib import Path
import tempfile
import unittest
from xml.etree import ElementTree as ET

from pptx import Presentation

from scripts.build_pipeline import build_parser, run_build
from scripts.svg_slot_binding import set_image_slot_href
from scripts.semantic_template_renderer import SemanticTemplate, render_slide
from scripts.slide_compiler import SlideCompileError
from scripts.template_compiler import compile_template


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates/layouts/literature_minimal"
LOGO = ROOT / "assets/logo.png"
SVG_NS = "http://www.w3.org/2000/svg"
XLINK_NS = "http://www.w3.org/1999/xlink"


def walk_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:
            yield from walk_shapes(shape.shapes)


def build_case(work, shell, logo=None):
    work.mkdir()
    payload = {slot["slot_id"]: "Check" for slot in shell["slots"] if slot.get("required", True)}
    if "CHAPTER_NUM" in payload:
        payload["CHAPTER_NUM"] = "01"
    if logo is not None:
        payload["LOGO_IMAGE"] = str(logo)
    plan = {
        "template_id": "literature_minimal",
        "production_scheme": "direct_editable",
        "slides": [{"page": "P01", "role": shell["role"], "shell_id": shell["shell_id"], "shell_payload": payload}],
    }
    plan_path = work / "plan.json"
    plan_path.write_text(json.dumps(plan), encoding="utf-8")
    args = build_parser().parse_args([
        str(plan_path), "--no-visual-plan", "--no-imagegen", "--no-svg",
        "--out-dir", str(work / "build"), "--pptx-out", str(work / "native.pptx"),
    ])
    return run_build(args)


class OptionalImageSlotTests(unittest.TestCase):
    def test_literature_logo_export_and_edit_roundtrip(self):
        ir = compile_template(TEMPLATE)["template_ir"]
        shells = {shell["shell_id"]: shell for shell in ir["shells"]}
        with tempfile.TemporaryDirectory() as tmp:
            for shell_id in ("toc", "chapter"):
                shell = shells[shell_id]
                logo_contract = next(slot for slot in shell["slots"] if slot["slot_id"] == "LOGO_IMAGE")
                self.assertEqual(logo_contract["kind"], "image")
                self.assertFalse(logo_contract["required"])
                source = ET.parse(ROOT / shell["svg_path"]).getroot()
                original_group = next(node for node in source.iter() if node.get("data-slot") == "LOGO_IMAGE")
                for with_logo in (False, True):
                    with self.subTest(shell=shell_id, with_logo=with_logo):
                        work = Path(tmp) / f"{shell_id}_{with_logo}"
                        report = build_case(work, shell, LOGO if with_logo else None)
                        self.assertEqual(report["status"], "pass")
                        svg = Path(report["pptx_render"]["svg_render"]["svg_files"][0])
                        self.assertNotIn("{{", svg.read_text(encoding="utf-8"))
                        root = ET.parse(svg).getroot()
                        groups = [node for node in root.iter() if node.get("data-slot") == "LOGO_IMAGE"]
                        self.assertEqual(len(groups), int(with_logo))
                        if with_logo:
                            group = groups[0]
                            self.assertIsNone(group.get("href"))
                            self.assertEqual(group.find(f"{{{SVG_NS}}}rect").attrib, original_group.find(f"{{{SVG_NS}}}rect").attrib)
                            image = group.find(f"{{{SVG_NS}}}image")
                            original_image = original_group.find(f"{{{SVG_NS}}}image")
                            for key in ("x", "y", "width", "height", "preserveAspectRatio"):
                                self.assertEqual(image.get(key), original_image.get(key))
                            self.assertEqual(image.get("href"), image.get(f"{{{XLINK_NS}}}href"))
                            self.assertTrue((svg.parent / image.get("href")).is_file())
                        else:
                            self.assertFalse(any(node.get("id") == f"{shell_id}-logo" for node in root.iter()))
                        deck = Presentation(work / "native.pptx")
                        shapes = list(walk_shapes(deck.slides[0].shapes))
                        pictures = [shape for shape in shapes if shape.shape_type == 13]
                        self.assertEqual(len(pictures), int(with_logo))
                        if pictures:
                            self.assertEqual(pictures[0].image.blob, LOGO.read_bytes())
                        textboxes = [shape for shape in shapes if shape.has_text_frame and shape.text.strip()]
                        self.assertTrue(textboxes)
                        textboxes[0].text = "Edited native text"
                        edited = work / "edited.pptx"
                        deck.save(edited)
                        reloaded = Presentation(edited)
                        self.assertIn("Edited native text", [shape.text for shape in walk_shapes(reloaded.slides[0].shapes) if shape.has_text_frame])

    def test_missing_logo_file_is_rejected(self):
        ir = compile_template(TEMPLATE)["template_ir"]
        with tempfile.TemporaryDirectory() as tmp:
            for shell in ir["shells"]:
                if shell["shell_id"] not in ("toc", "chapter"):
                    continue
                with self.subTest(shell=shell["shell_id"]):
                    work = Path(tmp) / shell["shell_id"]
                    with self.assertRaisesRegex(SlideCompileError, "references missing file"):
                        build_case(work, shell, Path(tmp) / "missing.png")
                    self.assertFalse((work / "native.pptx").exists())

    def test_cover_and_ending_preserve_grouped_titles_and_inline_metadata(self):
        ir = compile_template(TEMPLATE)["template_ir"]
        with tempfile.TemporaryDirectory() as tmp:
            for shell in ir["shells"]:
                if shell["shell_id"] not in ("cover", "ending"):
                    continue
                with self.subTest(shell=shell["shell_id"]):
                    work = Path(tmp) / shell["shell_id"]
                    report = build_case(work, shell)
                    self.assertEqual(report["status"], "pass")
                    svg = Path(report["pptx_render"]["svg_render"]["svg_files"][0])
                    root = ET.parse(svg).getroot()
                    self.assertNotIn("{{", svg.read_text(encoding="utf-8"))
                    for slot_id in ("AUTHOR", "DATE"):
                        span = next(node for node in root.iter() if node.get("data-slot") == slot_id)
                        self.assertEqual(span.text, "Check")
                        self.assertEqual(list(span), [])
                        self.assertIsNone(span.get("x"))
                        self.assertIsNone(span.get("y"))
                    texts = [shape.text for shape in walk_shapes(Presentation(work / "native.pptx").slides[0].shapes) if shape.has_text_frame]
                    self.assertIn("汇报人：Check", texts)
                    self.assertIn("日期：Check", texts)

    def test_semantic_renderer_binds_direct_and_grouped_optional_images(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            template = SemanticTemplate(root=work, template_id="fixture", layouts={}, variants=[])
            for grouped in (False, True):
                image = '<image x="10" y="20" width="194" height="40" href="{{LOGO}}"/>'
                slot = f'<g data-slot="LOGO"><rect width="214" height="56"/>{image}</g>' if grouped else image.replace('<image ', '<image data-slot="LOGO" ')
                (work / "slide.svg").write_text(f'<svg xmlns="{SVG_NS}">{slot}</svg>', encoding="utf-8")
                layout = {"layout_id": "fixture", "svg": "slide.svg", "slots": [{"slot_id": "LOGO", "kind": "image", "required": False}]}
                for with_logo in (False, True):
                    with self.subTest(grouped=grouped, with_logo=with_logo):
                        payload = {"LOGO": str(LOGO)} if with_logo else {}
                        tree, _ = render_slide(template, layout, {"slot_payload": payload}, plan_dir=work, assets_dir=work / "assets")
                        root = tree.getroot()
                        images = list(root.iter(f"{{{SVG_NS}}}image"))
                        self.assertEqual(len(images), int(with_logo))
                        self.assertEqual(len(list(root.iter(f"{{{SVG_NS}}}rect"))), int(grouped and with_logo))
                        if images:
                            self.assertTrue((work / images[0].get("href")).is_file())
                            self.assertEqual(images[0].get("width"), "194")
                            self.assertEqual(images[0].get("height"), "40")
                            self.assertNotIn("{{", ET.tostring(root, encoding="unicode"))

    def test_ambiguous_or_empty_image_groups_are_rejected_without_mutation(self):
        for children in ("<rect/>", "<image/><image/>"):
            with self.subTest(children=children):
                group = ET.fromstring(f'<g xmlns="{SVG_NS}" data-slot="LOGO">{children}</g>')
                before = ET.tostring(group)
                with self.assertRaisesRegex(ValueError, "exactly one image"):
                    set_image_slot_href(group, "logo.png")
                self.assertEqual(ET.tostring(group), before)

    def test_semantic_renderer_preserves_text_groups_and_inline_labels(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            (work / "slide.svg").write_text(
                f'<svg xmlns="{SVG_NS}"><g data-slot="TITLE"><rect width="200" height="50"/>'
                '<text x="10" y="30" font-size="24"><tspan>{{TITLE}}</tspan></text></g>'
                '<text x="20" y="80">Author: <tspan data-slot="AUTHOR">{{AUTHOR}}</tspan></text></svg>',
                encoding="utf-8",
            )
            template = SemanticTemplate(root=work, template_id="fixture", layouts={}, variants=[])
            layout = {"layout_id": "fixture", "svg": "slide.svg", "slots": [
                {"slot_id": slot_id, "kind": "text", "required": True} for slot_id in ("TITLE", "AUTHOR")
            ]}
            tree, _ = render_slide(template, layout, {"slot_payload": {"TITLE": "Check", "AUTHOR": "Ada"}}, plan_dir=work, assets_dir=work / "assets")
            root = tree.getroot()
            group = root.find(f"{{{SVG_NS}}}g")
            self.assertIsNotNone(group.find(f"{{{SVG_NS}}}rect"))
            self.assertEqual(group.find(f"{{{SVG_NS}}}text").get("x"), "10")
            self.assertEqual(["".join(node.itertext()) for node in root.iter(f"{{{SVG_NS}}}text")], ["Check", "Author: Ada"])
            span = next(node for node in root.iter() if node.get("data-slot") == "AUTHOR")
            self.assertEqual(list(span), [])
            self.assertIsNone(span.get("x"))


if __name__ == "__main__":
    unittest.main()
