"""Artifact-level smoke check and editable-object demonstration for test01."""
import json
from pathlib import Path
import sys
import unittest

from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from redraw_full_figure import check, PX

OUTPUT=ROOT/'output/bitmap_trace_test01/full_semantic'


@unittest.skipUnless(OUTPUT.is_dir(), "Optional local full-figure acceptance artifacts are not installed")
class FullFigureTests(unittest.TestCase):
    def test_native_scene_and_edit_roundtrip(self):
        data=json.loads((ROOT/'output/bitmap_trace_test01/editable_v2/annotations.json').read_text(encoding='utf-8'))
        report=check(OUTPUT/'test01_full_semantic.pptx',data)
        self.assertEqual(report['native_gradients'],3)
        prs=Presentation(OUTPUT/'test01_full_semantic.pptx')
        slide=prs.slides[0]
        shapes={s.name:s for s in slide.shapes}
        tree=shapes['Tree_036']
        tree_x=tree.left
        tree.left+=90*PX
        factory=shapes['Industrial_plant']
        factory_x=factory.left
        factory.left-=55*PX
        label=next(s for s in shapes['Industry_label'].shapes if s.name=='Text_industry')
        label.text_frame.paragraphs[0].runs[0].text='Editable text'
        destination=OUTPUT/'test01_edit_demo.pptx'
        prs.save(destination)
        loaded=Presentation(destination).slides[0]
        moved={s.name:s for s in loaded.shapes}
        self.assertEqual(moved['Tree_036'].left,tree_x+90*PX)
        self.assertEqual(moved['Industrial_plant'].left,factory_x-55*PX)
        self.assertEqual(len(moved['Tree_036']._element.xpath('.//a:path')),1)
        self.assertIn('Editable text',loaded._element.xpath('.//a:t/text()'))
        self.assertFalse(loaded._element.xpath('.//p:pic'))


if __name__=='__main__':
    unittest.main()
