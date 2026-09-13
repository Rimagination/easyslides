import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


def inventory_for(*, contract=None, element=None):
    from scripts.alignment_contract import default_contract

    return {
        "schema_version": "easyslides.slide_image_inventory.v1",
        "production_scheme": "image_full_rebuild",
        "reconstruction_mode": "preserve_complex_images",
        "alignment_contract": contract or default_contract(),
        "slides": [
            {
                "slide_id": "s01",
                "elements": [element or {
                    "element_id": "title",
                    "layer": "C",
                    "text": "Title",
                    "bbox_percent": {"x": 10, "y": 10, "w": 30, "h": 10},
                }],
                "source_text_lines": ["Title"],
                "completeness_check": {"performed": True, "layer_a_count": 0},
            }
        ],
    }


class AlignmentContractTests(unittest.TestCase):
    def test_native_table_columns_use_grid_width_and_keep_multiline_cells(self):
        from scripts.alignment_contract import _iter_pptx_text_frames, validate_pptx_alignment

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "table.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            table = slide.shapes.add_table(1, 2, Inches(1), Inches(.75), Inches(5), Inches(.75)).table
            table.columns[0].width = Inches(2)
            table.columns[1].width = Inches(3)
            table.cell(0, 0).text = "First"
            table.cell(0, 1).text = "Two\nlines"
            prs.save(path)
            frames = _iter_pptx_text_frames(path)[0]
            self.assertAlmostEqual(frames[0]["box"].width, 192)
            self.assertAlmostEqual(frames[1]["box"].x, 288)
            self.assertAlmostEqual(frames[1]["box"].width, 288)
            inventory = inventory_for(element={
                "element_id": "cell", "layer": "C", "text": "Two\nlines",
                "bbox_percent": {"x": 30, "y": 10, "w": 30, "h": 10},
            })
            self.assertEqual(validate_pptx_alignment(path, inventory)["status"], "pass")
            table.columns[0].width = Inches(2.5)
            table.columns[1].width = Inches(2.5)
            prs.save(path)
            report = validate_pptx_alignment(path, inventory)
            self.assertIn("ALIGNMENT-PPTX-BOX-DRIFT", {i["code"] for i in report["issues"]})

    def test_slide_relationship_targets_resolve_from_package_or_part(self):
        import io
        import zipfile
        from xml.etree import ElementTree as ET
        from scripts.alignment_contract import _pptx_slide_order

        buffer = io.BytesIO()
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(buffer)
        for target in ("slides/slide1.xml", "/ppt/slides/slide1.xml", "../ppt/slides/slide1.xml"):
            with self.subTest(target=target), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "relationships.pptx"
                with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(path, "w") as output:
                    for item in source.infolist():
                        data = source.read(item.filename)
                        if item.filename == "ppt/_rels/presentation.xml.rels":
                            root = ET.fromstring(data)
                            for relation in root:
                                if relation.get("Type", "").endswith("/slide"):
                                    relation.set("Target", target)
                            data = ET.tostring(root)
                        output.writestr(item, data)
                self.assertEqual(len(Presentation(path).slides), 1)
                self.assertEqual(_pptx_slide_order(path), ["ppt/slides/slide1.xml"])

    def test_image_inventory_requires_contract(self):
        from scripts.alignment_contract import validate_inventory_alignment

        report = validate_inventory_alignment(
            {
                "production_scheme": "image_full_rebuild",
                "slides": [{"slide_id": "s01", "elements": []}],
            }
        )

        self.assertEqual(report["status"], "fail")
        self.assertIn("ALIGNMENT-CONTRACT-MISSING", {item["code"] for item in report["issues"]})

    def test_svg_parent_transform_is_blocked(self):
        from scripts.alignment_contract import default_contract, validate_svg_alignment

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "slide.svg"
            path.write_text(
                """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720">
  <g transform="translate(80 0)">
    <text data-pptx-textbox="true" data-pptx-box-x="160" data-pptx-box-y="60" data-pptx-box-w="320" data-pptx-box-h="72" x="320" y="96" text-anchor="middle" data-pptx-valign="middle">Title</text>
  </g>
</svg>""",
                encoding="utf-8",
            )
            report = validate_svg_alignment([path], default_contract())

        self.assertEqual(report["status"], "fail")
        self.assertIn("ALIGNMENT-SVG-PARENT-TRANSFORM", {item["code"] for item in report["issues"]})

    def test_native_pptx_box_and_center_drift_are_blocking(self):
        from scripts.alignment_contract import validate_pptx_alignment

        inventory = inventory_for(
            element={
                "element_id": "title",
                "layer": "C",
                "text": "Title",
                "bbox_percent": {"x": 10, "y": 10, "w": 30, "h": 10},
                "alignment": {"center_lock": True, "vertical": "middle"},
            }
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "drifted.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(Inches(2), Inches(0.75), Inches(3), Inches(0.75)).text = "Title"
            prs.save(path)

            report = validate_pptx_alignment(path, inventory)

        codes = {item["code"] for item in report["issues"]}
        self.assertEqual(report["status"], "fail")
        self.assertIn("ALIGNMENT-PPTX-BOX-DRIFT", codes)
        self.assertIn("ALIGNMENT-PPTX-CENTER-DRIFT", codes)

    def test_native_pptx_matching_box_passes(self):
        from scripts.alignment_contract import validate_pptx_alignment

        inventory = inventory_for()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "aligned.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(0.75), Inches(3), Inches(0.75)).text = "Title"
            prs.save(path)

            report = validate_pptx_alignment(path, inventory)

        self.assertEqual(report["status"], "pass", report)
        self.assertEqual(report["blocking_count"], 0)


if __name__ == "__main__":
    unittest.main()
