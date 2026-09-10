"""Independent, annotation-assisted architectural redraw of the test01 crop.

Geometry is constructed from architectural planes and regular window grids;
no bitmap contour tracer or third-party project source is used here.
"""
from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
import xml.etree.ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

PX = 9525


def point(origin, vector, s, height=0, t=0):
    return [origin[0]+vector[0]*s, origin[1]+vector[1]*s+height*t]


def quad(origin, vector, height, s0, t0, s1, t1):
    if not (0 <= s0 < s1 <= 1 and 0 <= t0 < t1 <= 1):
        raise ValueError('Invalid facade interval')
    return [point(origin,vector,s0,height,t0), point(origin,vector,s1,height,t0),
            point(origin,vector,s1,height,t1), point(origin,vector,s0,height,t1)]


def primitive(name, points, color, role='face'):
    return {'name':name,'paths':[points],'fill':color,'role':role}


def is_lit(image, corners):
    cx = sum(p[0] for p in corners)/4
    cy = sum(p[1] for p in corners)/4
    tile = image.crop((round(cx)-1,round(cy)-1,round(cx)+2,round(cy)+2))
    return sum(r>205 and g>190 and r-b>35
               for r,g,b in (tile.getpixel((x,y)) for y in range(3) for x in range(3))) >= 3


def window_grid(spec, grid, image):
    cols, rows = grid['cols'], grid['rows']
    if not (isinstance(cols,int) and isinstance(rows,int) and 1<=cols<=100 and 1<=rows<=100):
        raise ValueError('Invalid grid count')
    origin = spec['origin'] if grid['side']=='left' else point(spec['origin'],spec['u'],1)
    vector = spec['u'] if grid['side']=='left' else spec['v']
    top, bottom, inset = .035, grid.get('bottom',.08), .045
    du, dv = (1-2*inset)/cols, (1-top-bottom)/rows
    gap = grid.get('gap',.23)
    if not (0 <= gap < 1 and 0 < top+bottom < 1):
        raise ValueError('Invalid grid spacing')
    windows = []
    for row in range(rows):
        for col in range(cols):
            s0, s1 = inset+du*(col+gap/2), inset+du*(col+1-gap/2)
            t0, t1 = top+dv*(row+.13), top+dv*(row+.87)
            corners = quad(origin,vector,spec['height'],s0,t0,s1,t1)
            color = 'FFF0B5' if is_lit(image,corners) else grid['glass']
            window = primitive(f"{spec['name']}_{grid['side']}_Window_R{row+1:02}_C{col+1:02}",corners,color,'window')
            window['facade'] = [origin,vector,spec['height']]
            window['cell'] = [s0,t0,s1,t1]
            if grid.get('pane_grid'):
                # Four regular panes share ONE compound shape: drag a whole window.
                sm, tm = (s0+s1)/2, (t0+t1)/2
                gu, gv = (s1-s0)*.045, (t1-t0)*.035
                window['paths'] = [quad(origin,vector,spec['height'],a,c,b,d)
                    for a,b in ((s0,sm-gu),(sm+gu,s1))
                    for c,d in ((t0,tm-gv),(tm+gv,t1))]
            windows.append(window)
    return {'name':f"{spec['name']}_{grid['side']}_Windows",'children':windows}


def make_block(spec, image):
    origin, u, v, h = spec['origin'],spec['u'],spec['v'],spec['height']
    if h<=0 or u[0]<=0 or v[0]<=0 or abs(u[0]*v[1]-u[1]*v[0])<1:
        raise ValueError(f"Degenerate building: {spec['name']}")
    corner = point(origin,u,1)
    rear = point(origin,v,1)
    right = point(corner,v,1)
    left_face = quad(origin,u,h,0,0,1,1)
    if spec.get('left_drop'):
        left_face[3][1] += spec['left_drop']
    children = [primitive(spec['name']+'_Left_wall',left_face,spec['walls'][0]),
                primitive(spec['name']+'_Right_wall',quad(corner,v,h,0,0,1,1),spec['walls'][1]),
                primitive(spec['name']+'_Roof',[origin,corner,right,rear],spec['roof'])]
    if color := spec.get('roof_inset'):
        # Uniform physical-width parapet rather than pixel-following roof outline.
        a, b = 2.2/math.hypot(*u), 2.2/math.hypot(*v)
        roof = [[origin[0]+u[0]*s+v[0]*t,origin[1]+u[1]*s+v[1]*t]
                for s,t in ((a,b),(1-a,b),(1-a,1-b),(a,1-b))]
        children.append(primitive(spec['name']+'_Roof_inset',roof,color))
    for side, start, vec in (('left',origin,u),('right',corner,v)):
        if spec.get('bands'):
            children.append(primitive(f"{spec['name']}_{side}_Plinth",
                quad(start,vec,h,0,.81,1,1),'C4CDD9','band'))
        for index in range(1,spec.get('bands',0)):
            t = index/spec['bands']
            children.append(primitive(f"{spec['name']}_{side}_Floor_band_{index}",
                quad(start,vec,h,0,t-.006,1,t+.006),'CAD3DF','band'))
    for grid in spec.get('grids',[]):
        children.append(window_grid(spec,grid,image))
    if door := spec.get('door'):
        start, vec = (origin,u) if door['side']=='left' else (corner,v)
        children.append(primitive(spec['name']+'_Door',quad(start,vec,h,door['u0'],door['top'],door['u1'],1),'62738F'))
    return {'name':spec['name'],'children':children}


def make_scene(model, source):
    return [{'name':building['name'], 'children':[make_block(s,source) for s in building['blocks']]}
            for building in model['buildings']]


def leaves(nodes):
    for node in nodes:
        if 'children' in node:
            yield from leaves(node['children'])
        else:
            yield node


def paint(shapes, nodes, scale=1, dx=0, dy=0):
    for node in nodes:
        if 'children' in node:
            group = shapes.add_group_shape()
            group.name = node['name']
            paint(group.shapes,node['children'],scale,dx,dy)
            continue
        # Milli-pixel local coordinates preserve the affine grid when enlarged.
        paths = [[(round((x*scale+dx)*1000),round((y*scale+dy)*1000)) for x,y in p] for p in node['paths']]
        builder = shapes.build_freeform(*paths[0][0],scale=PX/1000)
        for i, path in enumerate(paths):
            if i:
                builder.move_to(*path[0])
            builder.add_line_segments(path[1:],close=True)
        shape = builder.convert_to_shape()
        shape.name = node['name']
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(node['fill'])
        shape.line.fill.background()
        shape.shadow.inherit = False


def text(slide, content, x,y,w,h,size=22,color='29384D'):
    box = slide.shapes.add_textbox(Emu(x*PX),Emu(y*PX),Emu(w*PX),Emu(h*PX))
    box.text_frame.margin_left=box.text_frame.margin_top=0
    box.text_frame.margin_bottom=box.text_frame.margin_right=0
    box.text_frame.word_wrap=False
    run = box.text_frame.paragraphs[0].add_run()
    run.text=content
    run.font.name='Microsoft YaHei'
    run.font.size=Pt(size*.75)
    run.font.color.rgb=RGBColor.from_string(color)


def write_svg(scene, crop, path):
    ET.register_namespace('','http://www.w3.org/2000/svg')
    ns='{http://www.w3.org/2000/svg}'
    x,y,r,b=crop
    root=ET.Element(ns+'svg',{'viewBox':f'{x} {y} {r-x} {b-y}','width':str(r-x),'height':str(b-y)})
    def append(parent,nodes):
        for node in nodes:
            if 'children' in node:
                append(ET.SubElement(parent,ns+'g',{'id':node['name']}),node['children'])
            else:
                d=' '.join('M '+' L '.join(f'{a:.4f} {b:.4f}' for a,b in p)+' Z' for p in node['paths'])
                ET.SubElement(parent,ns+'path',{'id':node['name'],'fill':'#'+node['fill'],'d':d})
    append(root,scene)
    ET.ElementTree(root).write(path,encoding='utf-8',xml_declaration=True)


def save_decks(scene,model,source,output):
    crop=model['source_crop']
    source.crop(crop).save(output/'source_crop.png')
    prs=Presentation()
    prs.slide_width,prs.slide_height=Emu(1280*PX),Emu(720*PX)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    text(slide,'建筑局部 · 结构化重绘',48,32,1184,48,34)
    text(slide,'原图局部（位图参照）',48,110,550,40,22)
    text(slide,'规则立面与窗格（可编辑形状）',680,110,550,40,22)
    scale=550/(crop[2]-crop[0])
    slide.shapes.add_picture(str(output/'source_crop.png'),Emu(48*PX),Emu(186*PX),width=Emu(550*PX))
    paint(slide.shapes,scene,scale,680-crop[0]*scale,186-crop[1]*scale)
    text(slide,'窗口四边形  /  统一行列间距  /  平行边约束  /  按建筑分组',48,555,1190,36,23)
    text(slide,'仅重绘这一局部；窗格数量与遮挡后的结构按图像校准，未声称恢复原始设计文件。',48,624,1190,35,17,'66758A')
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    text(slide,'原生编辑页',48,26,1180,45,30)
    text(slide,'单击整栋移动；进入组后，可单独修改立面、窗户、屋顶。',48,77,1180,35,19,'66758A')
    paint(slide.shapes,scene,2.2,68-crop[0]*2.2,112-crop[1]*2.2)
    slide.notes_slide.notes_text_frame.text='Five named building groups. All architectural artwork is native editable geometry. No source bitmap on this slide.'
    # Keep a separate clean single-slide artifact with only the five building groups.
    clean=Presentation()
    clean.slide_width,clean.slide_height=prs.slide_width,prs.slide_height
    target=clean.slides.add_slide(clean.slide_layouts[6])
    paint(target.shapes,scene,2.2,68-crop[0]*2.2,40-crop[1]*2.2)
    clean.save(output/'buildings_editable.pptx')
    prs.save(output/'building_comparison.pptx')
    # Move one whole building and recolor one complete window in a real file.
    edited=Presentation(output/'buildings_editable.pptx')
    tower=next(s for s in edited.slides[0].shapes if s.name=='Front_office')
    tower.left-=Emu(48*PX)
    tower.top+=Emu(38*PX)
    def walk(shapes):
        for shape in shapes:
            if shape.shape_type==6:
                yield from walk(shape.shapes)
            else:
                yield shape
    selected=next(s for s in walk(tower.shapes) if s.name.endswith('left_Window_R01_C01'))
    selected.fill.fore_color.rgb=RGBColor.from_string('76B4D1')
    text(edited.slides[0],'操作示例：整栋移动 + 单窗改色',48,22,1180,40,24)
    edited.save(output/'building_move_demo.pptx')


def verify(scene,output):
    windows=[n for n in leaves(scene) if n['role']=='window']
    for window in windows:
        for p in window['paths']:
            assert len(p)==4
            assert abs(p[0][0]+p[2][0]-p[1][0]-p[3][0])<1e-7
            assert abs(p[0][1]+p[2][1]-p[1][1]-p[3][1])<1e-7
    prs=Presentation(output/'buildings_editable.pptx')
    slide=prs.slides[0]
    assert len(slide.shapes)==len(scene)==5
    assert not slide._element.xpath('.//p:pic')
    assert not slide._element.xpath('.//a:cubicBezTo')
    native=slide._element.xpath('.//p:sp')
    window_shapes=[s for s in native if '_Window_' in s.xpath('./p:nvSpPr/p:cNvPr')[0].get('name','')]
    assert len(window_shapes)==len(windows)
    for shape in window_shapes:
        moves=len(shape.xpath('.//a:moveTo'))
        assert len(shape.xpath('.//a:lnTo'))==moves*3
        assert len(shape.xpath('.//a:close'))==moves
        assert len(shape.xpath('./p:spPr/a:effectLst'))==1
    report={'building_groups':len(scene),'window_objects':len(windows),'native_primitives':len(native),
            'all_window_panes_have_four_vertices':True,'window_affine_parallelism':True,
            'embedded_images_on_editable_page':0,
            'scope':'Manually calibrated local architectural redraw; other figure regions unchanged.'}
    (output/'verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    return report


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,required=True)
    parser.add_argument('--model',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    source=Image.open(args.source).convert('RGB')
    model=json.loads(args.model.read_text(encoding='utf-8'))
    scene=make_scene(model,source)
    args.output.mkdir(parents=True,exist_ok=True)
    write_svg(scene,model['source_crop'],args.output/'buildings.svg')
    save_decks(scene,model,source,args.output)
    print(json.dumps(verify(scene,args.output)))


if __name__=='__main__':
    main()
