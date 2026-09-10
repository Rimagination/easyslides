"""Full annotation-assisted semantic redraw of test01, without bitmap tracing.

Scene geometry is explicitly drawn from visual understanding; repeated windows
reuse the project's architectural model. Source pixels only inform window light
colors. No third-party project source or generated raster artwork is embedded.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
import json
import math
from pathlib import Path
import time
import xml.etree.ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.util import Emu

from redraw_building_sample import make_scene, make_block
from rebuild_bitmap_figure import add_label
from svg_to_pptx.drawingml_converter import convert_svg_to_slide_shapes

NS='http://www.w3.org/2000/svg'
PX=9525
W,H=1792,1388
ET.register_namespace('',NS)


def element(parent,tag,**attrs):
    return ET.SubElement(parent,'{'+NS+'}'+tag,{k.replace('_','-'):str(v) for k,v in attrs.items()})


def group(parent,name):
    return element(parent,'g',id=name)


def path(parent,d,fill='none',stroke='none',width=1,**attrs):
    return element(parent,'path',d=d,fill=fill,stroke=stroke,stroke_width=width,**attrs)


def polygon(parent,points,fill,**attrs):
    return path(parent,'M '+' L '.join(f'{x:.3f} {y:.3f}' for x,y in points)+' Z',fill,**attrs)


def circle(parent,x,y,r,fill,stroke='none',width=1):
    return element(parent,'circle',cx=x,cy=y,r=r,fill=fill,stroke=stroke,stroke_width=width)


def gradient(defs,name,colors):
    grad=element(defs,'linearGradient',id=name,x1='0%',y1='0%',x2='0%',y2='100%')
    for offset,color in colors:
        element(grad,'stop',offset=f'{offset}%',stop_color=color)


def terrain(root):
    g=group(root,'Subsurface_block')
    polygon(g,[(0,497),(1064,955),(1064,1387),(0,929)],'#a3b9c2',stroke='#839dad',width=1.3)
    polygon(g,[(1064,955),(1791,544),(1791,920),(1064,1387)],'#a3b9c2',stroke='#839dad',width=1.3)
    g=group(root,'Water_depth')
    polygon(g,[(1064,955),(1791,544),(1791,834),(1064,1288)],'url(#water_depth)')
    path(g,'M850 872 C895 891 930 919 940 978 C948 1040 957 1085 996 1118 C1022 1139 1024 1137 1029 1186 C1035 1236 1041 1262 1064 1288 L1064 955 Z','url(#water_depth)')
    g=group(root,'Land_and_mountains')
    path(g,'M0 497 C130 462 226 397 319 334 C338 308 329 274 354 255 C372 241 402 232 419 218 C436 206 428 180 451 161 C478 139 503 116 519 96 C535 82 548 78 560 70 L596 57 Q607 52 619 65 C648 94 680 118 707 130 Q727 133 751 108 L767 99 L792 107 Q805 109 813 121 C830 133 850 133 866 153 Q877 173 894 176 C919 179 941 201 965 219 C993 241 1008 259 1029 273 C1073 301 1100 317 1128 322 C1161 330 1179 350 1209 362 L1595 480 L1791 544 L1064 955 Z','url(#mountain)')
    path(g,'M554 99 Q568 80 582 99 Q595 96 605 109 L590 105 Q580 109 574 103 Q565 97 554 99 Z','#fff')
    path(g,'M445 213 C448 202 454 177 464 175 C475 170 465 198 475 195 L491 191 Q499 164 505 178 L520 191 Q533 198 543 197 L543 204 L518 206 Q501 202 493 208 L476 215 L460 211 Z','#8c9e9f')
    path(g,'M550 196 Q575 184 582 169 Q591 166 598 173 Q613 182 630 182 L620 189 L594 193 L571 196 Z','#8c9d9d')
    path(g,'M651 178 Q666 170 661 154 L655 133 Q659 123 669 119 L686 128 L696 151 L685 160 L679 144 L671 132 L667 135 Q677 151 679 163 Z','#8b9c9f')
    path(g,'M386 306 C393 282 414 286 427 270 C437 252 460 266 482 275 C525 296 546 312 575 339 C598 359 608 371 618 393 L566 432 L422 408 Z','#cbd2b0',opacity=.68)
    # A closed shoreline ribbon tapers to zero at the upper land boundary.
    coast='M1595 480 C1535 478 1529 505 1480 517 C1432 530 1390 532 1388 551 C1370 568 1297 580 1268 608 L1235 640 C1208 656 1183 689 1172 726 C1163 761 1180 786 1119 792 C1030 790 981 825 851 872'
    path(g,coast+' L838 867 C969 810 1027 778 1118 779 C1162 777 1150 757 1159 723 C1170 682 1198 647 1227 630 L1259 598 C1291 568 1361 557 1376 541 C1382 520 1425 517 1477 504 C1521 493 1536 465 1560 472 L1595 480 Z','#fff0b8')
    g=group(root,'Sea_surface')
    path(g,coast+' L1064 955 L1791 544 Z','#89c7e7')
    path(g,'M851 872 L1064 955 L1791 544','none','#b6e5f5',1.5)
    path(g,'M1064 955 L1064 1288','none','#b6e5f5',1.3)
    g=group(root,'Water_table_boundary')
    path(g,'M0 497 L850 872 C892 891 930 917 940 978 C720 880 490 782 272 688 C160 644 75 608 0 557 Z','#d7cabe')
    path(g,'M0 557 C75 608 160 644 272 688 C490 782 720 880 940 978','none','#5699bb',1.9,stroke_dasharray='6 6')


def water_and_field(root):
    g=group(root,'River_network')
    path(g,'M714 130 C689 158 740 159 762 179 C783 199 748 226 706 236 C656 251 626 267 603 286 C577 310 600 322 636 332 C664 340 695 331 713 347 C732 370 697 392 671 412 C639 437 611 461 608 480 C605 495 626 501 646 496 C734 471 813 440 907 425 C961 416 1008 421 1033 442 C1056 459 1063 470 1043 491 C1019 519 966 537 925 565 C896 584 871 606 866 616 C862 629 900 629 937 628 L1044 625 C1110 618 1156 613 1197 615 C1230 615 1260 610 1286 592 L1300 618 C1264 631 1230 637 1175 633 C1119 632 1087 647 1038 648 L941 648 C886 648 846 645 842 625 C837 607 865 585 901 560 C947 525 994 505 1014 482 C1031 463 1021 452 995 444 C972 438 944 436 914 444 C824 458 734 492 654 516 C610 528 585 515 587 490 C590 458 624 433 657 406 C688 382 713 362 695 352 C682 344 653 344 631 339 C584 329 567 307 599 277 C625 255 657 243 704 229 C748 219 778 197 760 182 C743 166 690 162 707 139 Z','#89c8e8')
    path(g,'M1034 448 C1056 445 1062 422 1090 417 C1131 409 1152 408 1172 427 C1181 436 1184 450 1190 460 C1169 460 1147 455 1125 459 C1086 461 1070 451 1045 453 Z','#89c8e8')
    for d in [
        'M651 672 C680 648 733 668 759 661 C795 650 827 674 833 696 C838 716 789 730 748 734 C704 740 637 736 631 716 C624 703 637 682 651 672 Z',
        'M723 542 C731 531 758 529 776 535 C797 540 814 549 800 554 C780 560 754 558 732 557 C720 556 718 550 723 542 Z',
        'M802 511 C819 505 838 512 856 514 C879 519 870 526 851 526 L816 526 C796 524 791 517 802 511 Z',
        'M860 488 C877 491 896 484 916 484 C931 484 932 499 916 503 C894 508 871 502 859 500 C849 497 852 490 860 488 Z']:
        path(g,d,'#89c8e8')
    field=group(root,'Agricultural_field')
    path(field,'M88 532 C235 478 326 371 434 314 C484 348 566 372 629 388 C545 475 446 571 364 640 C267 611 168 574 88 532 Z','#efd978')
    for i in range(39):
        path(field,f'M{101+i*6.65:.2f} {534+i*2.43:.2f} C{249+i*5.5:.2f} {482+i*1.3:.2f} {345+i*4.9:.2f} {365+i*1.1:.2f} {435+i*4.75:.2f} {325+i*1.66:.2f}',
             'none','#cdb958',2.65,stroke_linecap='round')
    path(field,'M409 334 L434 316 C484 348 566 372 628 389 L582 431 C528 402 460 392 409 334 Z','#fff0b6',opacity=.33)


TREES=[
 (87,508,31,76),(133,488,44,87),(181,453,38,70),(247,395,27,66),(280,410,26,49),(307,361,27,45),(321,340,22,40),
 (500,279,25,61),(527,281,29,74),(546,268,24,54),(580,262,21,40),(598,258,21,43),(612,245,26,60),
 (640,203,23,41),(654,242,25,49),(700,224,25,48),(680,270,26,65),(556,312,34,70),(578,308,32,79),
 (622,299,30,60),(644,312,28,58),(670,326,32,62),(711,322,40,76),(730,278,30,75),
 (750,259,24,66),(767,263,28,58),(789,253,27,64),(744,296,24,47),(756,335,42,73),
 (799,305,30,68),(826,321,32,60),(831,282,25,48),(847,282,32,72),(883,268,24,58),(904,272,25,51),
 (888,398,35,84),(983,372,31,63),(1002,358,27,70),(1009,396,31,60),(1199,421,44,91),
 (668,467,41,79),(733,449,25,62),(756,445,28,58),(950,475,38,65),(768,522,32,59),
 (731,537,25,52),(831,554,28,60),(779,636,35,69),(811,639,45,94)]


def tree_outline(x,base,w,h,variant=0):
    top=base-h
    parts=[f'M{x:.3f} {top:.3f}']
    # Each side has intentional branch tiers; no pixel-boundary noise.
    for side in (1,-1):
        levels=range(1,10) if side==1 else range(9,0,-1)
        if side==-1:
            parts += [f'Q{x+w*.42:.3f} {base-h*.075:.3f} {x+w*.07:.3f} {base-h*.09:.3f}',
                      f'L{x+w*.065:.3f} {base:.3f} L{x-w*.065:.3f} {base:.3f} L{x-w*.07:.3f} {base-h*.09:.3f}',
                      f'Q{x-w*.42:.3f} {base-h*.075:.3f} {x-w*.46:.3f} {base-h*.14:.3f}']
        for level in levels:
            t=level/10
            broad=w*(.04+.45*t)*(1+variant*.025)
            px=x+side*broad
            py=top+h*t*.91
            if side==1:
                parts.append(f'Q{x+side*broad*.73:.3f} {py-h*.09:.3f} {px:.3f} {py:.3f} L{x+side*broad*.72:.3f} {py-h*.01:.3f}')
            else:
                parts.append(f'L{px:.3f} {py:.3f} Q{x+side*broad*.75:.3f} {py-h*.04:.3f} {x+side*broad*.70:.3f} {py-h*.075:.3f}')
    parts.append('Z')
    return ' '.join(parts)


def trees(root):
    colors=['#417e49','#609362','#729c70']
    # Back-to-front order is explicit by ground-contact position.
    for i,(x,y,w,h) in sorted(enumerate(TREES,1),key=lambda item:item[1][1]):
        g=group(root,f'Tree_{i:03d}')
        shade=2 if i in (4,7,13,15,36,40,42,44,47) else (1 if i%3==0 else 0)
        path(g,tree_outline(x,y,w,h,i%3-1),colors[shade])


def add_architecture(parent,nodes,factor):
    for node in nodes:
        if 'children' in node:
            add_architecture(group(parent,node['name']),node['children'],factor)
        else:
            d=' '.join('M '+' L '.join(f'{x*factor:.4f} {y*factor:.4f}' for x,y in p)+' Z' for p in node['paths'])
            path(parent,d,'#'+node['fill'],id=node['name'])


def building_spec(name,origin,u,v,height,cols=7,rows=7,bands=0):
    spec={'name':name,'origin':origin,'u':u,'v':v,'height':height,
          'walls':['93A0B4','CAD3DF'],'roof':'D9E0E8'}
    if bands:
        spec.update(bands=bands,roof_inset='8998AF')
    else:
        spec['grids']=[{'side':'left','cols':cols,'rows':rows,'bottom':.09,'glass':'72839E'},
                       {'side':'right','cols':max(3,round(cols*v[0]/u[0])),'rows':rows,'bottom':.09,'glass':'8294AF'}]
    return spec


def buildings(root,source,model):
    f=W/source.width
    add_architecture(root,make_scene(model,source),f)
    near=[
      ('Town_high_rise',[building_spec('High_rise',[1030,578],[58,20],[40,-14],112,bands=7)]),
      ('Town_rear_annex',[building_spec('Town_annex',[925,681],[25,9],[22,-8],64,4,7)]),
      ('Town_main_office',[building_spec('Town_main',[809,684],[51,18],[61,-22],102,8,12),
                           building_spec('Town_upper',[859,630],[36,13],[28,-10],52,5,6)]),
      ('Town_front_office',[building_spec('Town_front',[880,748],[42,15],[72,-24],50,5,5)]),
      ('Town_right_offices',[building_spec('Town_right_rear',[1010,684],[67,23],[59,-21],65,8,7),
                             building_spec('Town_right_front',[1010,712],[34,12],[37,-13],65,4,6)])]
    for name,specs in near:
        parent=group(root,name)
        for original in specs:
            spec=deepcopy(original)
            for key in ('origin','u','v'):
                spec[key]=[n/f for n in spec[key]]
            spec['height']/=f
            add_architecture(parent,[make_block(spec,source)],f)


def factory(root):
    g=group(root,'Industrial_plant')
    polygon(g,[(351,617),(481,576),(646,607),(495,659),(417,670)],'#809744')
    for x,y,w,h in [(436,505,45,74),(504,481,46,68),(572,461,25,93)]:
        tube=group(g,f'Plant_tank_{x}')
        path(tube,f'M{x} {y} Q{x+w/2} {y-12} {x+w} {y} L{x+w} {y+h} Q{x+w/2} {y+h+10} {x} {y+h} Z','#929db1')
        path(tube,f'M{x+w*.53} {y} L{x+w} {y} L{x+w} {y+h} Q{x+w*.8} {y+h+5} {x+w*.53} {y+h+5} Z','#7f8ca3')
        element(tube,'ellipse',cx=x+w/2,cy=y,rx=w/2,ry=6,fill='#c4cfdc',stroke='#65758e',stroke_width=1.2)
        if w<30:
            for dy in (24,47,70):
                path(tube,f'M{x+1} {y+dy} Q{x+w/2} {y+dy+6} {x+w-1} {y+dy}','none','#c5cfdb',1.2)
    polygon(g,[(425,570),(495,593),(495,659),(425,635)],'#8492aa')
    polygon(g,[(495,593),(646,553),(646,607),(495,659)],'#929db1')
    polygon(g,[(425,570),(576,530),(646,553),(495,593)],'#d9dfe7')
    for i in range(3):
        x=437+i*18
        polygon(g,[(x,583+i*6),(x+10,586+i*6),(x+10,617+i*6),(x,614+i*6)],'#bac7d8')
    for i,x in enumerate((505,521,557,571)):
        top=599-(x-505)*.27
        polygon(g,[(x,top),(x+10,top-3),(x+10,top+25),(x,top+28)],'#c1ccda')
    polygon(g,[(539,624),(549,621),(549,641),(539,645)],'#475e7a')
    polygon(g,[(449,565),(517,547),(565,562),(497,581)],'#b8c3d3')
    polygon(g,[(449,565),(497,581),(497,585),(449,569)],'#8191a9')
    for x,y in ((490,563),(518,554)):
        element(g,'ellipse',cx=x,cy=y,rx=14,ry=5.5,fill='#74849c',stroke='#536b85',stroke_width=1)
        for angle in (0,120,240):
            a=math.radians(angle)
            path(g,f'M{x} {y} L{x+11*math.cos(a):.2f} {y+4*math.sin(a):.2f}','none','#3e5673',1.5)


def rotate_points(points,x,y,angle):
    c,s=math.cos(math.radians(angle)),math.sin(math.radians(angle))
    return [(x+a*c-b*s,y+a*s+b*c) for a,b in points]


def arrow_badge(root,name,x,y,color,direction=0):
    g=group(root,name)
    circle(g,x,y,18,color)
    polygon(g,rotate_points([(-2,11),(-2,-2),(-7,1),(0,-13),(7,1),(2,-2),(2,11)],x,y,direction),'#fff')


def fish(root,name,x,y,w,h,angle=0):
    g=group(root,name)
    pts=rotate_points([(-w*.5,0),(-w*.72,-h*.6),(-w*.63,0),(-w*.72,h*.6),(-w*.5,0)],x,y,angle)
    polygon(g,pts,'#4b9ac3')
    # The body and fins share the fish group.
    d=f'M{x-w*.5} {y} C{x-w*.3} {y-h} {x+w*.4} {y-h*.7} {x+w*.5} {y} C{x+w*.3} {y+h*.55} {x-w*.3} {y+h*.5} {x-w*.5} {y} Z'
    path(g,d,'#4b9ac3')


def biological_details(root):
    g=group(root,'Groundwater_microbiomes')
    circle(g,138,815,85,'#89c8e8','#526f83',1.5)
    pale='#f0edc2'
    # A diatom with curved silica striae.
    diatom=group(g,'Microbe_diatom')
    element(diatom,'ellipse',cx=116,cy=770,rx=16,ry=9,fill='none',stroke=pale,stroke_width=1.6)
    for i in range(27):
        a=2*math.pi*i/27
        x,y=116+17*math.cos(a),770+10*math.sin(a)
        xx,yy=116+29*math.cos(a+.15),770+23*math.sin(a+.15)
        path(diatom,f'M{x:.2f} {y:.2f} Q{xx-3:.2f} {yy-3:.2f} {xx:.2f} {yy:.2f}','none',pale,1.1)
    star=group(g,'Microbe_star')
    circle(star,176,782,3,'none',pale,1.3)
    for i in range(8):
        a=i*math.pi/4
        x,y=176+24*math.cos(a),782+24*math.sin(a)
        path(star,f'M176 782 L{x:.2f} {y:.2f}','none',pale,1.4)
        circle(star,x,y,1.5,pale)
    cell=group(g,'Microbe_flagellate')
    path(cell,'M126 820 C104 825 142 869 153 860 C163 851 137 817 126 820 Z','none',pale,1.6)
    path(cell,'M128 820 C124 831 142 853 152 861 M130 820 C142 833 142 850 151 859 M130 820 C119 810 122 803 113 799 M133 822 C134 807 145 805 146 800 M148 856 C160 858 161 871 165 873','none',pale,1.2)
    for i,(x,y,r) in enumerate([(178,824,9),(190,825,8),(199,827,7),(184,837,8),(194,839,7),(186,847,6)]):
        cluster=group(g,f'Microbe_coccus_{i}')
        pts=[(x+r*math.cos(a*math.pi/3),y+r*.8*math.sin(a*math.pi/3)) for a in range(6)]
        polygon(cluster,pts,'none',stroke=pale,width=1.1)
    path(g,'M83 851 C116 865 132 892 117 881 C101 869 90 860 80 858 M99 838 C120 850 122 875 111 875 C94 872 88 862 80 859 M113 870 C120 875 127 872 133 867','none',pale,1.6)
    path(g,'M102 824 C112 820 113 809 124 812 C127 819 119 821 120 829 M123 814 C121 803 129 800 133 794','none',pale,1.4)
    g=group(root,'Biochemical_cycle')
    circle(g,407,930,87,'#fff','#405b76',3.4)
    polygon(g,[(488,921),(500,916),(495,931)],'#405b76')
    polygon(g,[(334,881),(336,866),(346,876)],'#405b76')
    for x,y,color,edge in [(347,871,'#f6d77b','#ba9940'),(468,871,'#f6b7b5','#c96b70'),(347,989,'#afe0f1','#68abc7'),(465,990,'#ace3df','#4fbdb7')]:
        circle(g,x,y,14,color,edge,1.3)
    element(g,'rect',x=371,y=894,width=67,height=25,fill='#bfc5d0')
    element(g,'rect',x=356,y=935,width=99,height=26,fill='#d3e7c5')
    g=group(root,'Key_functions_arrow')
    path(g,'M220 861 L309 896','none','#405b76',3.2)
    polygon(g,[(304,888),(315,899),(300,897)],'#405b76')
    g=group(root,'Water_health_microbiomes')
    circle(g,837,1118,86,'#fff','#667474',1.4)
    bacteria=[
        ('M786 1069 C808 1056 801 1087 824 1080',7),
        ('M774 1085 C793 1097 801 1083 814 1089',8),
        ('M815 1106 C816 1087 849 1093 853 1082 C854 1070 873 1072 873 1058',9),
        ('M809 1131 C818 1109 838 1116 842 1127',9),
        ('M770 1111 C770 1124 783 1137 788 1130',10),
        ('M795 1118 C800 1132 794 1140 790 1145',9),
        ('M867 1125 C877 1129 883 1137 894 1139',11),
        ('M896 1099 C899 1110 900 1120 903 1129',10),
        ('M812 1168 C825 1168 833 1170 842 1172',10),
        ('M850 1192 C856 1170 878 1183 883 1169 C887 1156 903 1162 907 1147',9)]
    for i,(d,width) in enumerate(bacteria,1):
        b=group(g,f'Bacterium_{i:02d}')
        path(b,d,'none','#cb696b',width+2,stroke_linecap='round',stroke_linejoin='round')
        path(b,d,'none','#efa0a1',width-1,stroke_linecap='round',stroke_linejoin='round')
    for i,(x,y,r) in enumerate([(833,1050,8),(875,1106,8),(834,1144,8),(852,1146,8),(862,1156,8)]):
        circle(group(g,f'Bacterial_coccus_{i}'),x,y,r,'#f6b5b4','#cf7173',1.6)
    for i in range(4):
        path(g,f'M{814+i*4} {1101+i*2} C{842+i*2} {1098+i*5} {829+i*3} {1117+i*2} {853+i*2} {1115+i*3}','none','#d77c7e',1)
        path(g,f'M{777+i*2} {1155+i*4} C790 {1152+i*4} 786 {1165+i*3} 807 {1163+i*3}','none','#d77c7e',1)
    for i in range(3):
        path(g,f'M{891+i*3} 1093 C{895+i*3} 1082 {879+i*4} 1076 {884+i*3} 1063','none','#d77c7e',1)


def labels_and_badges(root):
    for name,x,y,w,h in [('Irrigation_label',62,571,253,35),('Industry_label',315,656,190,34),
      ('Exchange_label',501,728,190,34),('Drinking_label',685,814,174,34),('Water_exchange_label',874,1228,134,62)]:
        element(group(root,name),'rect',x=x,y=y,width=w,height=h,fill='#fff',opacity=.63)
    g=group(root,'Legend')
    element(g,'rect',x=1327,y=167,width=462,height=124,fill='#fff',stroke='#666',stroke_width=1.3)
    arrow_badge(g,'Legend_subsurface_arrow',1362,203,'#db6868')
    arrow_badge(g,'Legend_surface_arrow',1362,255,'#4f9dca')
    for name,x,y,color,angle in [('Irrigation_input',188,633,'#4f9dca',0),('Industrial_input',410,716,'#4f9dca',0),
      ('Anthropogenic_input',410,775,'#db6868',180),('Surface_exchange',578,792,'#4f9dca',0),
      ('Subsurface_exchange',618,807,'#db6868',180),('Drinking_input',772,874,'#4f9dca',0),
      ('Geogenic_input',404,1095,'#db6868',0),('Water_exchange_out',1009,1238,'#4f9dca',90),('Water_exchange_in',1009,1283,'#db6868',-90)]:
        arrow_badge(root,name,x,y,color,angle)
    g=group(root,'Ion_panel')
    element(g,'rect',x=558,y=925,width=148,height=205,fill='#fff')
    for i,(color,edge) in enumerate([('#fbb9b7','#d16a6d'),('#dac1df','#b18abb'),('#97dbf7','#5ab0d3'),('#ffbf80','#eba460'),
                                      ('#f6d980','#d2af58'),('#bfdfa8','#8cbd77'),('#a2e0e0','#49b5bd'),('#d6dde5','#99a6b8')]):
        x,y=568+(i%2)*69,939+(i//2)*47
        element(g,'rect',x=x,y=y,width=57,height=35,rx=6,ry=6,fill=color,stroke=edge,stroke_width=1.3)
    g=group(root,'Annotation_leaders')
    for x,y1,y2 in [(24,578,623),(135,900,1020),(404,1114,1138),(632,1130,1241),(837,1205,1326)]:
        path(g,f'M{x} {y1} L{x} {y2}','none','#606665',1.3)


LABEL_GROUPS={
 'legend_subsurface':'Legend','legend_surface':'Legend',
 'irrigation':'Irrigation_label','industry':'Industry_label','exchange_land':'Exchange_label','drinking':'Drinking_label','exchange_water':'Water_exchange_label',
 **{k:'Biochemical_cycle' for k in ('metal','organics','cycle_c','cycle_n','cycle_s','cycle_p')},
 **{k:'Ion_panel' for k in ('vanadium','manganese','uranium','arsenic','chromium','iron','organic_carbon','nitrate')}}


def create_art(source,architecture):
    root=ET.Element('{'+NS+'}svg',{'width':str(W),'height':str(H),'viewBox':f'0 0 {W} {H}'})
    defs=element(root,'defs')
    gradient(defs,'mountain',[(0,'#9aa5b8'),(15,'#a2b09a'),(30,'#a6ba54'),(100,'#a5b952')])
    gradient(defs,'water_depth',[(0,'#88c8e8'),(40,'#6eaec7'),(100,'#408b98')])
    terrain(root)
    water_and_field(root)
    # Ground-contact shadows are semantic objects, not per-window effects.
    shadows=group(root,'Building_ground_shadows')
    for pts in [[(1071,530),(1225,482),(1301,571),(1193,598)],
                [(714,793),(861,746),(994,787),(893,833)],[(978,748),(1082,710),(1150,762),(1081,798)]]:
        polygon(shadows,pts,'#809744')
    trees(root)
    buildings(root,source,architecture)
    factory(root)
    for i,(x,y,w,h) in enumerate([(1635,635,34,7),(1628,673,32,9),(1651,698,34,7),(1701,650,35,8)]):
        fish(root,f'Surface_fish_{i}',x,y,w,h)
    deep=group(root,'Deep_water_organisms')
    for i,(x,y,w,h) in enumerate([(1458,833,40,11),(1480,823,32,13),(1500,825,19,11)]):
        fish(deep,f'Deep_fish_{i}',x,y,w,h)
    biological_details(root)
    labels_and_badges(root)
    return root


def make_ppt(art_path,annotations,output):
    xml,media,rels,targets=convert_svg_to_slide_shapes(art_path)
    assert not media and not rels
    parsed=parse_xml(xml.encode('utf-8'))
    prs=Presentation()
    prs.slide_width,prs.slide_height=Emu(W*PX),Emu(H*PX)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    for child in list(parsed.xpath('./p:cSld/p:spTree')[0])[2:]:
        slide.shapes._spTree.insert_element_before(deepcopy(child),'p:extLst')
    names=dict(targets)
    for shape in slide.shapes:
        if shape.shape_id in names:
            shape.name=names[shape.shape_id]
    groups={s.name:s for s in slide.shapes if s.shape_type==6}
    for item in annotations['labels']:
        parent=groups[LABEL_GROUPS[item['id']]] if item['id'] in LABEL_GROUPS else slide
        add_label(parent,item,annotations,(W,H))
    # Nested text addition can otherwise reuse shape IDs from the root tree.
    for i,node in enumerate(slide._element.xpath('.//p:cNvPr'),1):
        node.set('id',str(i))
    slide.notes_slide.notes_text_frame.text=(
        'Whole-figure semantic redraw from visual interpretation. Native shapes and text only. '
        'Buildings use regular facade grids; trees, arrows and biological insets are grouped. '
        'Silhouettes, tree arrangement and microbe morphology are approximate. '
        'Occluded geometry is inferred. This is not recovery of the source design file.')
    target=output/'test01_full_semantic.pptx'
    prs.save(target)
    return target


def svg_with_text(root,data,target):
    root=deepcopy(root)
    index={n.get('id'):n for n in root.iter() if n.get('id')}
    for item in data['labels']:
        parent=index[LABEL_GROUPS[item['id']]] if item['id'] in LABEL_GROUPS else root
        x,y,x2,y2=item['box']
        size=item.get('font_size',22)
        if item.get('align')=='center':
            x=(x+x2)/2
        for i,line in enumerate(item['text'].split('\n')):
            txt=element(parent,'text',x=x,y=y+size*.80+i*size*1.16,font_family='Arial',font_size=size,fill='#383838',
                        text_anchor='middle' if item.get('align')=='center' else 'start')
            if 'runs' not in item:
                txt.text=line
            else:
                for content,baseline in item['runs']:
                    span=element(txt,'tspan',font_size=size*(.65 if baseline else 1),
                                 baseline_shift='super' if baseline>0 else ('sub' if baseline<0 else 'baseline'))
                    span.text=content
    ET.ElementTree(root).write(target,encoding='utf-8',xml_declaration=True)


def check(ppt,data):
    slide=Presentation(ppt).slides[0]
    assert not slide._element.xpath('.//p:pic')
    text={n.xpath('./p:nvSpPr/p:cNvPr')[0].get('name'):''.join(n.xpath('.//a:t/text()'))
          for n in slide._element.xpath('.//p:sp') if n.xpath('.//a:t')}
    for label in data['labels']:
        assert text['Text_'+label['id']]==label['text'].replace('\n',''),label['id']
    identifiers=[n.get('id') for n in slide._element.xpath('.//p:cNvPr')]
    assert len(identifiers)==len(set(identifiers))
    tree_shapes=[s for s in slide.shapes if s.name.startswith('Tree_')]
    assert len(tree_shapes)==len(TREES)
    for tree in tree_shapes:
        assert len(tree._element.xpath('.//a:path'))==1
        assert len(tree._element.xpath('.//a:close'))==1
    return {'top_level_objects':len(slide.shapes),'tree_objects':len(tree_shapes),
            'native_text_blocks':len(text),'native_gradients':len(slide._element.xpath('.//a:gradFill')),
            'embedded_pictures':0,'native_shapes':len(slide._element.xpath('.//p:sp')),
            'scope':'Whole-figure redraw; geometry and biological detail approximated from visual interpretation.'}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,required=True)
    parser.add_argument('--architecture',type=Path,required=True)
    parser.add_argument('--annotations',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    started=time.perf_counter()
    source=Image.open(args.source).convert('RGB')
    model=json.loads(args.architecture.read_text(encoding='utf-8'))
    annotations=json.loads(args.annotations.read_text(encoding='utf-8'))
    args.output.mkdir(parents=True,exist_ok=True)
    root=create_art(source,model)
    art=args.output/'geometry.svg'
    ET.ElementTree(root).write(art,encoding='utf-8',xml_declaration=True)
    svg_with_text(root,annotations,args.output/'test01_full_semantic.svg')
    ppt=make_ppt(art,annotations,args.output)
    report=check(ppt,annotations)
    report['generation_seconds']=round(time.perf_counter()-started,3)
    (args.output/'verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(report))


if __name__=='__main__':
    main()
