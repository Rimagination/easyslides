#!/usr/bin/env python3
"""Shared geometry contract for image-to-editable slide reconstruction.

The image reconstruction route has three coordinate representations:
source-image percentages, SVG canvas pixels, and native PPTX geometry.  This
module keeps their invariants in one place so a visual preview cannot hide a
shift that appears after native export.
"""

from __future__ import annotations

from copy import deepcopy
import math
import posixpath
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET
import zipfile

try:
    from scripts import layout_metrics
except ImportError:  # pragma: no cover
    import layout_metrics


CONTRACT_SCHEMA_VERSION = "easyslides.alignment_contract.v1"
REPORT_SCHEMA_VERSION = "easyslides.alignment_contract_report.v1"
CANVAS_WIDTH = 1280.0
CANVAS_HEIGHT = 720.0
IDENTITY_MATRIX = layout_metrics.IDENTITY_MATRIX
PPTX_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def default_contract() -> dict[str, Any]:
    """Return the default contract for newly initialized image projects."""
    return {
        "schema_version": CONTRACT_SCHEMA_VERSION,
        "coordinate_space": "source_percent_to_1280x720",
        "canvas": {"width": CANVAS_WIDTH, "height": CANVAS_HEIGHT},
        "geometry_policy": {
            "native_text_box": "absolute_canvas",
            "parent_transform": "forbidden_for_editable_text",
            "match_inventory_boxes": True,
            "box_tolerance_px": 4.0,
            "center_lock_tolerance_px": 2.0,
        },
        "protected_regions": [],
    }


def _issue(
    code: str,
    severity: str,
    message: str,
    *,
    slide_id: str | None = None,
    slide_number: int | None = None,
    element_id: str | None = None,
    details: dict[str, Any] | None = None,
) -> dict[str, Any]:
    payload: dict[str, Any] = {"code": code, "severity": severity, "message": message}
    if slide_id:
        payload["slide_id"] = slide_id
    if slide_number is not None:
        payload["slide_number"] = slide_number
    if element_id:
        payload["element_id"] = element_id
    if details:
        payload["details"] = details
    return payload


def _number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _box_from_mapping(payload: Any) -> layout_metrics.Box | None:
    if not isinstance(payload, dict):
        return None
    values = []
    for key in ("x", "y"):
        values.append(_number(payload.get(key)))
    width = payload.get("w", payload.get("width"))
    height = payload.get("h", payload.get("height"))
    values.extend((_number(width), _number(height)))
    if any(value is None for value in values):
        return None
    x, y, w, h = values  # type: ignore[misc]
    if w <= 0 or h <= 0:
        return None
    return layout_metrics.Box(x=x, y=y, width=w, height=h)


def _percent_box(payload: Any) -> layout_metrics.Box | None:
    box = _box_from_mapping(payload)
    if box is None:
        return None
    return layout_metrics.Box(
        x=box.x * CANVAS_WIDTH / 100.0,
        y=box.y * CANVAS_HEIGHT / 100.0,
        width=box.width * CANVAS_WIDTH / 100.0,
        height=box.height * CANVAS_HEIGHT / 100.0,
    )


def _box_payload(box: layout_metrics.Box) -> dict[str, float]:
    return {
        "x": round(box.x, 2),
        "y": round(box.y, 2),
        "width": round(box.width, 2),
        "height": round(box.height, 2),
        "cx": round(box.cx, 2),
        "cy": round(box.cy, 2),
    }


def _merge_contract(base: dict[str, Any], override: Any) -> dict[str, Any]:
    result = deepcopy(base)
    if not isinstance(override, dict):
        return result
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(result.get(key), dict):
            result[key] = _merge_contract(result[key], value)
        else:
            result[key] = deepcopy(value)
    return result


def contract_for_slide(inventory: dict[str, Any], slide: dict[str, Any]) -> dict[str, Any]:
    return _merge_contract(inventory.get("alignment_contract") or default_contract(), slide.get("alignment_contract"))


def _is_image_reconstruction_inventory(inventory: dict[str, Any]) -> bool:
    return str(inventory.get("production_scheme") or "").startswith("image_")


def _validate_region(region: Any, *, location: str) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    if not isinstance(region, dict):
        return [_issue("ALIGNMENT-REGION-MALFORMED", "blocking", f"{location} must be an object.")]
    if "bbox_percent" in region:
        box = _box_from_mapping(region.get("bbox_percent"))
        bounds = layout_metrics.Box(0, 0, 100, 100)
        if box is None or not layout_metrics.contains(bounds, box, tolerance=0.001):
            issues.append(_issue("ALIGNMENT-REGION-OUT-OF-BOUNDS", "blocking", f"{location}.bbox_percent must be a positive in-bounds percentage box."))
    else:
        box = _box_from_mapping(region)
        bounds = layout_metrics.Box(0, 0, CANVAS_WIDTH, CANVAS_HEIGHT)
        if box is None or not layout_metrics.contains(bounds, box, tolerance=0.001):
            issues.append(_issue("ALIGNMENT-REGION-OUT-OF-BOUNDS", "blocking", f"{location} must be a positive in-bounds 1280x720 box."))
    return issues


def validate_contract(contract: Any, *, required: bool = False) -> list[dict[str, Any]]:
    """Validate contract shape without coupling to the inventory schema."""
    if not isinstance(contract, dict):
        if not required:
            return []
        return [
            _issue(
                "ALIGNMENT-CONTRACT-MISSING",
                "blocking" if required else "warning",
                "Image reconstruction inventory must declare alignment_contract.",
            )
        ]

    issues: list[dict[str, Any]] = []
    if contract.get("schema_version") != CONTRACT_SCHEMA_VERSION:
        issues.append(
            _issue(
                "ALIGNMENT-CONTRACT-SCHEMA",
                "blocking",
                f"alignment_contract.schema_version must be {CONTRACT_SCHEMA_VERSION}.",
            )
        )
    if contract.get("coordinate_space") != "source_percent_to_1280x720":
        issues.append(
            _issue(
                "ALIGNMENT-CONTRACT-COORDINATE-SPACE",
                "blocking",
                "alignment_contract.coordinate_space must map source percentages to the 1280x720 canvas.",
            )
        )
    canvas = contract.get("canvas")
    if not isinstance(canvas, dict) or _number(canvas.get("width")) != CANVAS_WIDTH or _number(canvas.get("height")) != CANVAS_HEIGHT:
        issues.append(_issue("ALIGNMENT-CONTRACT-CANVAS", "blocking", "alignment_contract.canvas must be exactly 1280x720."))
    policy = contract.get("geometry_policy")
    if not isinstance(policy, dict):
        issues.append(_issue("ALIGNMENT-CONTRACT-POLICY", "blocking", "alignment_contract.geometry_policy is required."))
    else:
        if policy.get("native_text_box") != "absolute_canvas":
            issues.append(_issue("ALIGNMENT-CONTRACT-TEXT-SPACE", "blocking", "Editable text geometry must be declared in absolute canvas coordinates."))
        if policy.get("parent_transform") != "forbidden_for_editable_text":
            issues.append(_issue("ALIGNMENT-CONTRACT-PARENT-TRANSFORM", "blocking", "Editable text must not inherit a parent transform in the reconstruction route."))
        if policy.get("match_inventory_boxes") is not True:
            issues.append(_issue("ALIGNMENT-CONTRACT-BOX-MATCH", "blocking", "The reconstruction contract must compare native text frames with source inventory boxes."))
        for key in ("box_tolerance_px", "center_lock_tolerance_px"):
            value = _number(policy.get(key))
            if value is None or value < 0:
                issues.append(_issue("ALIGNMENT-CONTRACT-TOLERANCE", "blocking", f"geometry_policy.{key} must be a non-negative number."))
    regions = contract.get("protected_regions", [])
    if not isinstance(regions, list):
        issues.append(_issue("ALIGNMENT-CONTRACT-PROTECTED-REGIONS", "blocking", "protected_regions must be a list."))
    else:
        for index, region in enumerate(regions, start=1):
            issues.extend(_validate_region(region, location=f"protected_regions[{index}]"))
    return issues


def validate_inventory_alignment(inventory: dict[str, Any]) -> dict[str, Any]:
    """Validate alignment declarations before a PPTX is inspected."""
    required = _is_image_reconstruction_inventory(inventory)
    contract = inventory.get("alignment_contract")
    issues = validate_contract(contract, required=required)
    if not isinstance(contract, dict):
        contract = default_contract()

    slides = inventory.get("slides") if isinstance(inventory.get("slides"), list) else []
    checked_text_count = 0
    for slide_number, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            continue
        slide_id = str(slide.get("slide_id") or f"s{slide_number:02d}")
        slide_contract = contract_for_slide({"alignment_contract": contract}, slide)
        for index, region in enumerate(slide_contract.get("protected_regions", []), start=1):
            issues.extend(_validate_region(region, location=f"{slide_id}.alignment_contract.protected_regions[{index}]"))
        elements = slide.get("elements") if isinstance(slide.get("elements"), list) else []
        for element_index, element in enumerate(elements, start=1):
            if not isinstance(element, dict) or str(element.get("layer", "")).upper() != "C":
                continue
            checked_text_count += 1
            element_id = str(element.get("element_id") or f"{slide_id}_e{element_index:02d}")
            alignment = element.get("alignment") if isinstance(element.get("alignment"), dict) else {}
            geometry_space = str(
                alignment.get("geometry_space")
                or element.get("geometry_space")
                or "absolute_canvas"
            ).strip().lower()
            if geometry_space not in {"absolute_canvas", "canvas", "absolute"}:
                issues.append(
                    _issue(
                        "ALIGNMENT-INVENTORY-TEXT-SPACE",
                        "blocking",
                        "Layer C geometry must use absolute canvas coordinates after source scaling.",
                        slide_id=slide_id,
                        element_id=element_id,
                    )
                )
            parent_transform = alignment.get("parent_transform", element.get("parent_transform"))
            if parent_transform not in (None, "", False, "none", "identity", "flattened"):
                issues.append(
                    _issue(
                        "ALIGNMENT-INVENTORY-PARENT-TRANSFORM",
                        "blocking",
                        "Layer C text declares a parent transform; flatten the text box to absolute canvas geometry.",
                        slide_id=slide_id,
                        element_id=element_id,
                        details={"parent_transform": parent_transform},
                    )
                )
            if alignment.get("center_lock") is True:
                vertical = str(alignment.get("vertical") or alignment.get("valign") or "middle").lower()
                if vertical not in {"middle", "center", "ctr"}:
                    issues.append(
                        _issue(
                            "ALIGNMENT-INVENTORY-CENTER-VALIGN",
                            "blocking",
                            "A center-locked Layer C text element must declare middle vertical alignment.",
                            slide_id=slide_id,
                            element_id=element_id,
                        )
                    )

    blocking_count = sum(item["severity"] == "blocking" for item in issues)
    warning_count = sum(item["severity"] == "warning" for item in issues)
    return {
        "schema_version": REPORT_SCHEMA_VERSION,
        "contract_schema_version": contract.get("schema_version"),
        "status": "fail" if blocking_count else "pass",
        "blocking_count": blocking_count,
        "warning_count": warning_count,
        "slide_count": len(slides),
        "checked_text_count": checked_text_count,
        "issues": issues,
    }


def _normal_text(value: Any) -> str:
    return "".join(str(value or "").split())


def _shape_name(sp: ET.Element) -> str:
    c_nv_pr = sp.find("p:nvSpPr/p:cNvPr", PPTX_NS)
    return str(c_nv_pr.get("name") or "") if c_nv_pr is not None else ""


def _text_content(sp: ET.Element) -> str:
    return "".join(node.text or "" for node in sp.findall(".//a:t", PPTX_NS)).strip()


def _pptx_slide_order(pptx_path: Path) -> list[str]:
    with zipfile.ZipFile(pptx_path) as zf:
        presentation = ET.fromstring(zf.read("ppt/presentation.xml"))
        rels = ET.fromstring(zf.read("ppt/_rels/presentation.xml.rels"))
        relmap = {
            rel.get("Id"): rel.get("Target")
            for rel in rels.findall("rel:Relationship", PPTX_NS)
            if rel.get("Id") and rel.get("Target")
        }
        result: list[str] = []
        for slide_id in presentation.findall(".//p:sldId", PPTX_NS):
            target = relmap.get(slide_id.get(f"{{{PPTX_NS['r']}}}id"))
            if target:
                result.append(posixpath.normpath(posixpath.join("ppt", target)).lstrip("/"))
        return result


def _pptx_box(xfrm: ET.Element | None) -> layout_metrics.Box | None:
    if xfrm is None:
        return None
    off = xfrm.find("a:off", PPTX_NS)
    ext = xfrm.find("a:ext", PPTX_NS)
    if off is None or ext is None:
        return None
    return layout_metrics.Box(
        x=layout_metrics.emu_to_px(off.get("x")),
        y=layout_metrics.emu_to_px(off.get("y")),
        width=layout_metrics.emu_to_px(ext.get("cx")),
        height=layout_metrics.emu_to_px(ext.get("cy")),
    )


def _pptx_group_matrix(grp: ET.Element) -> layout_metrics.AffineMatrix:
    xfrm = grp.find("p:grpSpPr/a:xfrm", PPTX_NS)
    outer = _pptx_box(xfrm)
    ch_off = xfrm.find("a:chOff", PPTX_NS) if xfrm is not None else None
    ch_ext = xfrm.find("a:chExt", PPTX_NS) if xfrm is not None else None
    if outer is None or ch_off is None or ch_ext is None:
        return IDENTITY_MATRIX
    child_x = layout_metrics.emu_to_px(ch_off.get("x"))
    child_y = layout_metrics.emu_to_px(ch_off.get("y"))
    child_w = layout_metrics.emu_to_px(ch_ext.get("cx"))
    child_h = layout_metrics.emu_to_px(ch_ext.get("cy"))
    sx = outer.width / child_w if child_w else 1.0
    sy = outer.height / child_h if child_h else 1.0
    matrix = (
        sx,
        0.0,
        0.0,
        sy,
        outer.x - child_x * sx,
        outer.y - child_y * sy,
    )
    return layout_metrics.matrix_multiply(_pptx_rotation(xfrm, outer), matrix)


def _pptx_rotation(xfrm: ET.Element | None, box: layout_metrics.Box) -> layout_metrics.AffineMatrix:
    if xfrm is None:
        return IDENTITY_MATRIX
    raw = _number(xfrm.get("rot"))
    if raw is None or not raw:
        return IDENTITY_MATRIX
    return layout_metrics.rotate_matrix(raw / 60000.0, box.cx, box.cy)


def _iter_pptx_text_frames(pptx_path: Path) -> list[list[dict[str, Any]]]:
    frames: list[list[dict[str, Any]]] = []
    with zipfile.ZipFile(pptx_path) as zf:
        for slide_name in _pptx_slide_order(pptx_path):
            root = ET.fromstring(zf.read(slide_name))
            slide_frames: list[dict[str, Any]] = []

            def walk(parent: ET.Element, parent_matrix: layout_metrics.AffineMatrix) -> None:
                for child in parent:
                    local = child.tag.rsplit("}", 1)[-1]
                    if local == "grpSp":
                        walk(child, layout_metrics.matrix_multiply(parent_matrix, _pptx_group_matrix(child)))
                        continue
                    if local == "sp":
                        xfrm = child.find("p:spPr/a:xfrm", PPTX_NS)
                        box = _pptx_box(xfrm)
                        text = _text_content(child)
                        if box is not None and text:
                            box = layout_metrics.transform_box(box, _pptx_rotation(xfrm, box))
                            box = layout_metrics.transform_box(box, parent_matrix)
                            slide_frames.append({"text": text, "box": box, "shape_name": _shape_name(child), "kind": "text_box"})
                        continue
                    if local == "graphicFrame":
                        xfrm = child.find("p:xfrm", PPTX_NS)
                        table_box = _pptx_box(xfrm)
                        table = child.find(".//a:tbl", PPTX_NS)
                        if table_box is not None and table is not None:
                            columns = [layout_metrics.emu_to_px(col.get("cx")) for col in table.findall("a:tblGrid/a:gridCol", PPTX_NS)]
                            rows = [layout_metrics.emu_to_px(row.get("h")) for row in table.findall("a:tr", PPTX_NS)]
                            if columns and rows:
                                y = table_box.y
                                for row_index, (row, row_height) in enumerate(zip(table.findall("a:tr", PPTX_NS), rows)):
                                    x = table_box.x
                                    for col_index, (cell, col_width) in enumerate(zip(row.findall("a:tc", PPTX_NS), columns)):
                                        text = "".join(node.text or "" for node in cell.findall(".//a:t", PPTX_NS)).strip()
                                        if text:
                                            cell_box = layout_metrics.Box(x=x, y=y, width=col_width, height=row_height)
                                            cell_box = layout_metrics.transform_box(cell_box, parent_matrix)
                                            slide_frames.append({"text": text, "box": cell_box, "shape_name": f"{_shape_name(child)}[{row_index},{col_index}]", "kind": "table_cell"})
                                        x += col_width
                                    y += row_height
                        continue
                    walk(child, parent_matrix)

            walk(root, IDENTITY_MATRIX)
            frames.append(slide_frames)
    return frames


def _expected_text_box(element: dict[str, Any]) -> layout_metrics.Box | None:
    return _percent_box(element.get("bbox_percent"))


def _protected_regions(contract: dict[str, Any]) -> list[tuple[str, layout_metrics.Box]]:
    result: list[tuple[str, layout_metrics.Box]] = []
    for index, region in enumerate(contract.get("protected_regions", []), start=1):
        if not isinstance(region, dict):
            continue
        box = _percent_box(region.get("bbox_percent")) if "bbox_percent" in region else _box_from_mapping(region)
        if box is not None:
            result.append((str(region.get("id") or f"protected_{index:02d}"), box))
    return result


def validate_pptx_alignment(pptx_path: str | Path, inventory: dict[str, Any]) -> dict[str, Any]:
    """Compare native PPTX text-box frames with source inventory geometry."""
    path = Path(pptx_path)
    contract = inventory.get("alignment_contract")
    issues = validate_inventory_alignment(inventory).get("issues", [])
    if not isinstance(contract, dict):
        blocking_count = sum(item["severity"] == "blocking" for item in issues)
        return {
            "schema_version": REPORT_SCHEMA_VERSION,
            "contract_schema_version": None,
            "status": "fail" if blocking_count else "pass",
            "blocking_count": blocking_count,
            "warning_count": sum(item["severity"] == "warning" for item in issues),
            "checked_text_count": 0,
            "issues": issues,
        }

    try:
        from pptx import Presentation

        prs = Presentation(str(path))
        slide_width_px = layout_metrics.emu_to_px(prs.slide_width)
        slide_height_px = layout_metrics.emu_to_px(prs.slide_height)
    except Exception as exc:  # pragma: no cover - malformed PPTX is covered by the structural gate
        issues.append(_issue("ALIGNMENT-PPTX-READ", "blocking", f"Cannot read PPTX geometry: {exc}"))
        slide_width_px = CANVAS_WIDTH
        slide_height_px = CANVAS_HEIGHT

    frames_by_slide = _iter_pptx_text_frames(path)
    slides = inventory.get("slides") if isinstance(inventory.get("slides"), list) else []
    policy = contract.get("geometry_policy") if isinstance(contract.get("geometry_policy"), dict) else {}
    box_tolerance = _number(policy.get("box_tolerance_px"))
    center_tolerance = _number(policy.get("center_lock_tolerance_px"))
    box_tolerance = 4.0 if box_tolerance is None else box_tolerance
    center_tolerance = 2.0 if center_tolerance is None else center_tolerance
    scale_x = CANVAS_WIDTH / slide_width_px if slide_width_px else 1.0
    scale_y = CANVAS_HEIGHT / slide_height_px if slide_height_px else 1.0
    checked_text_count = 0

    for slide_number, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            continue
        expected_elements = [
            element
            for element in (slide.get("elements") if isinstance(slide.get("elements"), list) else [])
            if isinstance(element, dict) and str(element.get("layer", "")).upper() == "C"
        ]
        if not expected_elements:
            continue
        actual_frames = []
        if slide_number <= len(frames_by_slide):
            for frame in frames_by_slide[slide_number - 1]:
                box = frame["box"]
                actual_frames.append({
                    **frame,
                    "box": layout_metrics.Box(box.x * scale_x, box.y * scale_y, box.width * scale_x, box.height * scale_y),
                })
        used: set[int] = set()
        slide_contract = contract_for_slide(inventory, slide)
        protected = _protected_regions(slide_contract)
        slide_id = str(slide.get("slide_id") or f"s{slide_number:02d}")
        for element_index, element in enumerate(expected_elements, start=1):
            checked_text_count += 1
            element_id = str(element.get("element_id") or f"{slide_id}_e{element_index:02d}")
            expected_box = _expected_text_box(element)
            if expected_box is None:
                continue
            expected_text = _normal_text(element.get("text"))
            candidates = [
                (index, frame)
                for index, frame in enumerate(actual_frames)
                if index not in used and _normal_text(frame.get("text")) == expected_text
            ]
            if not candidates:
                issues.append(
                    _issue(
                        "ALIGNMENT-PPTX-TEXT-MISSING",
                        "blocking",
                        "Layer C text has no matching native PPTX text frame for geometry validation.",
                        slide_id=slide_id,
                        slide_number=slide_number,
                        element_id=element_id,
                        details={"text": str(element.get("text") or "")[:120], "expected_box": _box_payload(expected_box)},
                    )
                )
                continue
            chosen_index, chosen = min(
                candidates,
                key=lambda item: abs(item[1]["box"].cx - expected_box.cx) + abs(item[1]["box"].cy - expected_box.cy),
            )
            used.add(chosen_index)
            actual_box = chosen["box"]
            deltas = {
                key: abs(getattr(actual_box, key) - getattr(expected_box, key))
                for key in ("x", "y", "width", "height")
            }
            if max(deltas.values()) > box_tolerance:
                issues.append(
                    _issue(
                        "ALIGNMENT-PPTX-BOX-DRIFT",
                        "blocking",
                        "Native PPTX text frame drifted from the source inventory box.",
                        slide_id=slide_id,
                        slide_number=slide_number,
                        element_id=element_id,
                        details={
                            "text": str(element.get("text") or "")[:120],
                            "expected_box": _box_payload(expected_box),
                            "actual_box": _box_payload(actual_box),
                            "deltas": {key: round(value, 2) for key, value in deltas.items()},
                            "shape_name": chosen.get("shape_name"),
                        },
                    )
                )
            alignment = element.get("alignment") if isinstance(element.get("alignment"), dict) else {}
            center_lock = alignment.get("center_lock") is True or alignment.get("center_lock") is None and str(alignment.get("vertical") or "").lower() in {"middle", "center", "ctr"}
            if center_lock and (abs(actual_box.cx - expected_box.cx) > center_tolerance or abs(actual_box.cy - expected_box.cy) > center_tolerance):
                issues.append(
                    _issue(
                        "ALIGNMENT-PPTX-CENTER-DRIFT",
                        "blocking",
                        "Center-locked native text frame drifted from the source center.",
                        slide_id=slide_id,
                        slide_number=slide_number,
                        element_id=element_id,
                        details={
                            "expected_center": {"x": round(expected_box.cx, 2), "y": round(expected_box.cy, 2)},
                            "actual_center": {"x": round(actual_box.cx, 2), "y": round(actual_box.cy, 2)},
                            "tolerance_px": center_tolerance,
                        },
                    )
                )
            allowed_region = str(alignment.get("protected_region_id") or "")
            for region_id, region_box in protected:
                if layout_metrics.overlap_area(actual_box, region_box) <= 1.0:
                    continue
                if allowed_region == region_id and layout_metrics.contains(region_box, actual_box, tolerance=box_tolerance):
                    continue
                issues.append(
                    _issue(
                        "ALIGNMENT-PPTX-PROTECTED-OVERLAP",
                        "blocking",
                        f"Native text frame overlaps protected region {region_id} without an explicit chrome binding.",
                        slide_id=slide_id,
                        slide_number=slide_number,
                        element_id=element_id,
                        details={"text_box": _box_payload(actual_box), "protected_region": _box_payload(region_box)},
                    )
                )

    blocking_count = sum(item["severity"] == "blocking" for item in issues)
    warning_count = sum(item["severity"] == "warning" for item in issues)
    return {
        "schema_version": REPORT_SCHEMA_VERSION,
        "contract_schema_version": contract.get("schema_version"),
        "status": "fail" if blocking_count else "pass",
        "blocking_count": blocking_count,
        "warning_count": warning_count,
        "slide_count": len(slides),
        "checked_text_count": checked_text_count,
        "issues": issues,
    }


def _svg_text_issues(root: ET.Element, *, svg_name: str, contract: dict[str, Any]) -> list[dict[str, Any]]:
    policy = contract.get("geometry_policy") if isinstance(contract.get("geometry_policy"), dict) else {}
    canvas = layout_metrics.Box(0, 0, CANVAS_WIDTH, CANVAS_HEIGHT)
    issues: list[dict[str, Any]] = []
    index = 0

    def walk(elem: ET.Element, parent_matrix: layout_metrics.AffineMatrix) -> None:
        nonlocal index
        index += 1
        local = elem.tag.rsplit("}", 1)[-1]
        if local == "text" and str(elem.get("data-pptx-textbox") or "").lower() in {"1", "true", "yes"}:
            required = [elem.get(f"data-pptx-box-{key}") for key in ("x", "y", "w", "h")]
            declared = None
            parsed = [_number(value) for value in required]
            if all(value is not None for value in parsed):
                declared = layout_metrics.Box(*(value for value in parsed if value is not None))
            element_id = str(elem.get("data-reconstruction-element-id") or elem.get("data-slot-id") or f"text_{index:03d}")
            if declared is None or not layout_metrics.contains(canvas, declared, tolerance=box_tolerance(policy)):
                issues.append({"code": "ALIGNMENT-SVG-BOX-INVALID", "severity": "blocking", "svg_file": svg_name, "element_id": element_id, "message": "Editable SVG text must declare a positive in-bounds absolute canvas box."})
            if parent_matrix != IDENTITY_MATRIX:
                issues.append({"code": "ALIGNMENT-SVG-PARENT-TRANSFORM", "severity": "blocking", "svg_file": svg_name, "element_id": element_id, "message": "Editable SVG text inherits a parent transform; flatten it to absolute canvas geometry before PPTX export."})
            if declared is not None:
                valign = str(elem.get("data-pptx-valign") or "").lower()
                if not valign:
                    issues.append({"code": "ALIGNMENT-SVG-VALIGN-MISSING", "severity": "blocking", "svg_file": svg_name, "element_id": element_id, "message": "Editable SVG text must declare data-pptx-valign so native vertical placement is deterministic."})
                if str(elem.get("data-center-lock") or "").lower() in {"1", "true", "yes"} and valign not in {"middle", "center", "ctr"}:
                    issues.append({"code": "ALIGNMENT-SVG-CENTER-VALIGN", "severity": "blocking", "svg_file": svg_name, "element_id": element_id, "message": "A center-locked SVG text element must use middle vertical alignment."})
                anchor = str(elem.get("text-anchor") or "start").lower()
                x = layout_metrics.parse_float(elem.get("x"))
                expected_x = declared.cx if anchor == "middle" else declared.right if anchor == "end" else declared.x
                if abs(x - expected_x) > box_tolerance(policy):
                    issues.append({"code": "ALIGNMENT-SVG-ANCHOR-DRIFT", "severity": "blocking", "svg_file": svg_name, "element_id": element_id, "message": "SVG text anchor does not match its declared absolute text box.", "details": {"text_x": round(x, 2), "expected_x": round(expected_x, 2), "box": _box_payload(declared)}})
        matrix = layout_metrics._element_matrix(elem, parent_matrix)
        for child in elem:
            walk(child, matrix)

    walk(root, IDENTITY_MATRIX)
    return issues


def box_tolerance(policy: dict[str, Any]) -> float:
    value = _number(policy.get("box_tolerance_px"))
    return value if value is not None else 4.0


def validate_svg_alignment(paths: Iterable[str | Path], contract: dict[str, Any]) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    checked = 0
    for raw_path in paths:
        path = Path(raw_path)
        try:
            root = ET.parse(path).getroot()
        except (OSError, ET.ParseError) as exc:
            issues.append({"code": "ALIGNMENT-SVG-READ", "severity": "blocking", "svg_file": str(path), "message": f"Cannot read reconstruction SVG: {exc}"})
            continue
        checked += 1
        issues.extend(_svg_text_issues(root, svg_name=str(path), contract=contract))
    blocking_count = sum(item["severity"] == "blocking" for item in issues)
    return {
        "schema_version": REPORT_SCHEMA_VERSION,
        "contract_schema_version": contract.get("schema_version"),
        "status": "fail" if blocking_count else "pass",
        "blocking_count": blocking_count,
        "warning_count": sum(item["severity"] == "warning" for item in issues),
        "svg_count": checked,
        "issues": issues,
    }


__all__ = [
    "CONTRACT_SCHEMA_VERSION",
    "REPORT_SCHEMA_VERSION",
    "CANVAS_WIDTH",
    "CANVAS_HEIGHT",
    "default_contract",
    "contract_for_slide",
    "validate_contract",
    "validate_inventory_alignment",
    "validate_pptx_alignment",
    "validate_svg_alignment",
]
