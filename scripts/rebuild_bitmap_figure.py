"""Annotation-assisted raster figure reconstruction, using EasySlides' renderer.

Independently implemented for this project. No cell-series code is imported,
vendored or executed. Text is supplied as reviewed annotations, not inferred by
this script. Only explicitly annotated isolated trees are reconstructed.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
import json
from pathlib import Path
import subprocess
import sys
import xml.etree.ElementTree as ET

import cv2
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Emu, Pt

from svg_to_pptx.drawingml_converter import convert_svg_to_slide_shapes

SVG = "http://www.w3.org/2000/svg"
PX = 9525


def scaled_box(box, reference, size):
    values = [round(v * size[i % 2] / reference[i % 2]) for i, v in enumerate(box)]
    x1, y1, x2, y2 = values
    if not (0 <= x1 < x2 <= size[0] and 0 <= y1 < y2 <= size[1]):
        raise ValueError(f"Annotation outside image: {values}")
    return values


def validate_annotations(data, size):
    seen = set()
    for kind, reference in (("labels", data["reference_size"]), ("trees", data["tree_coordinate_size"])):
        for item in data[kind]:
            if item["id"] in seen:
                raise ValueError(f"Duplicate annotation ID: {item['id']}")
            seen.add(item["id"])
            scaled_box(item["box"], reference, size)
            if kind == "labels":
                for box in item.get("erase_boxes", []):
                    scaled_box(box, reference, size)
                if not item["text"].strip():
                    raise ValueError("Empty text annotation")
                if "runs" in item and "".join(r[0] for r in item["runs"]) != item["text"]:
                    raise ValueError(f"Rich-text mismatch: {item['id']}")
                if any(r[1] not in (-1, 0, 1) for r in item.get("runs", [])):
                    raise ValueError("Unsupported text baseline")


def prepare_layers(source, data, output):
    """Remove annotated glyphs/trees before tracing, avoiding layered ghosts."""
    height, width = source.shape[:2]
    size = (width, height)
    text_mask = np.zeros((height, width), np.uint8)
    gray = cv2.cvtColor(source, cv2.COLOR_BGR2GRAY)
    for item in data["labels"]:
        for box in item.get("erase_boxes", [item["box"]]):
            x1, y1, x2, y2 = scaled_box(box, data["reference_size"], size)
            # ponytail: dark text only; colored/light text needs an explicit glyph mask.
            glyph = (gray[y1:y2, x1:x2] < 145).astype(np.uint8) * 255
            glyph = cv2.dilate(glyph, np.ones((3, 3), np.uint8))
            text_mask[y1:y2, x1:x2] |= glyph
    cleaned = cv2.inpaint(source, text_mask, 4, cv2.INPAINT_TELEA)
    hsv = cv2.cvtColor(source, cv2.COLOR_BGR2HSV)
    tree_mask = np.zeros_like(text_mask)
    trees = []
    for item in data["trees"]:
        x1, y1, x2, y2 = scaled_box(item["box"], data["tree_coordinate_size"], size)
        tile = hsv[y1:y2, x1:x2]
        mask = cv2.inRange(tile, (44, 45, 0), (86, 255, 170))
        mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, np.ones((3, 3), np.uint8))
        contours, _ = cv2.findContours(mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        if not contours:
            raise ValueError(f"No tree contour: {item['id']}")
        contour = max(contours, key=cv2.contourArea)
        if cv2.contourArea(contour) < 60:
            raise ValueError(f"Insufficient tree area: {item['id']}")
        solid = np.zeros_like(mask)
        cv2.drawContours(solid, [contour], -1, 255, -1)
        color = np.median(source[y1:y2, x1:x2][solid > 0], axis=0).astype(int)
        simplified = cv2.approxPolyDP(contour, 0.65, True).reshape(-1, 2)
        simplified += (x1, y1)
        trees.append({"id":item["id"], "points":simplified.tolist(),
                      "fill":"#" + "".join(f"{v:02x}" for v in color[::-1])})
        tree_mask[y1:y2, x1:x2] |= solid
    erased = cv2.dilate(tree_mask, np.ones((7, 7), np.uint8))
    cleaned = cv2.inpaint(cleaned, erased, 6, cv2.INPAINT_TELEA)
    for name, pixels in (("background.png", cleaned), ("text_mask.png", text_mask), ("tree_mask.png", erased)):
        if not cv2.imwrite(str(output / name), pixels):
            raise OSError(f"Cannot write {name}")
    return trees


def trace_background(image, svg, python, packages):
    # Reuse the already installed VTracer in its matching Python ABI; no install.
    worker = ("import sys; sys.path.insert(0,sys.argv[1]); import vtracer; "
              "vtracer.convert_image_to_svg_py(sys.argv[2],sys.argv[3],"
              "colormode='color',hierarchical='stacked',mode='spline',"
              "filter_speckle=2,color_precision=8,layer_difference=6,"
              "corner_threshold=60,length_threshold=3.5,max_iterations=10,"
              "splice_threshold=45,path_precision=3)")
    subprocess.run([str(python), "-c", worker, str(packages), str(image), str(svg)], check=True)


def assemble_scene(background, trees, target, size, data=None):
    ET.register_namespace("", SVG)
    source = ET.parse(background).getroot()
    root = ET.Element(f"{{{SVG}}}svg", {"width":str(size[0]), "height":str(size[1]),
                                      "viewBox":f"0 0 {size[0]} {size[1]}"})
    group = ET.SubElement(root, f"{{{SVG}}}g", {"id":"traced_background"})
    # Preserve original path order and each compound path intact.
    for child in source:
        group.append(deepcopy(child))
    for item in (data or {}).get("rectangles", []):
        x1, y1, x2, y2 = scaled_box(item["box"], data["reference_size"], size)
        ET.SubElement(group, f"{{{SVG}}}rect", {"x":str(x1), "y":str(y1),
            "width":str(x2-x1), "height":str(y2-y1), "fill":"none", "stroke":"#666666",
            "stroke-width":str(size[0]/data["reference_size"][0])})
    for tree in trees:
        d = "M " + " L ".join(f"{x} {y}" for x, y in tree["points"]) + " Z"
        ET.SubElement(root, f"{{{SVG}}}path", {"id":tree["id"], "d":d, "fill":tree["fill"]})
    ET.ElementTree(root).write(target, encoding="utf-8", xml_declaration=True)


def add_label(slide, item, data, size):
    x1, y1, x2, y2 = scaled_box(item["box"], data["reference_size"], size)
    font_px = item.get("font_size", 22) * size[1] / data["reference_size"][1]
    shape = slide.shapes.add_textbox(Emu(x1 * PX), Emu(round((y1 - font_px * .16) * PX)),
                                   Emu((x2-x1+5) * PX), Emu(round((y2-y1+font_px*.3) * PX)))
    shape.name = "Text_" + item["id"]
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(item["text"].split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER if item.get("align") == "center" else PP_ALIGN.LEFT
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = Pt(font_px * .75 * 1.16)
        for content, baseline in item.get("runs", [[line, 0]]):
            run = paragraph.add_run()
            run.text = content
            run.font.name = data.get("font", "Arial")
            run.font.size = Pt(font_px * .75 * (.65 if baseline else 1))
            run.font.color.rgb = RGBColor(56, 56, 56)
            if baseline:
                run._r.get_or_add_rPr().set("baseline", "40000" if baseline > 0 else "-22000")
    return shape


def make_deck(scene, data, trees, size, output):
    xml, media, relationships, _ = convert_svg_to_slide_shapes(scene)
    if media or relationships:
        raise ValueError("Figure unexpectedly contains embedded media")
    converted = parse_xml(xml.encode("utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = (Emu(v * PX) for v in size)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    source_shapes = converted.xpath("./p:cSld/p:spTree")[0]
    for child in list(source_shapes)[2:]:
        slide.shapes._spTree.insert_element_before(deepcopy(child), "p:extLst")
    slide.shapes[0].name = "Background_details__ungroup_to_edit"
    for index, tree in enumerate(trees, 1):
        slide.shapes[index].name = tree["id"]
    for item in data["labels"]:
        add_label(slide, item, data, size)
    slide.notes_slide.notes_text_frame.text = (
        f"Annotation-assisted reconstruction. {len(data['labels'])} manually transcribed text blocks; "
        f"{len(trees)} selected tree silhouettes. Dense forest and buildings remain in the "
        "background detail group. Tree colors are simplified; hidden backgrounds "
        "are locally interpolated. Original bitmap is not embedded.")
    target = output / "test01_editable_v2.pptx"
    prs.save(target)
    # A separate real movement artifact, not a label claiming movement works.
    moved = Presentation(target)
    if demo := data.get("demo"):
        tree = next(s for s in moved.slides[0].shapes if s.name == demo["tree_id"])
        tree.left += Emu(demo["shift"][0] * PX)
        tree.top += Emu(demo["shift"][1] * PX)
        label = next(s for s in moved.slides[0].shapes if s.name == "Text_" + demo["text_id"])
        label.text_frame.paragraphs[0].runs[0].text = demo["replacement"]
        moved.save(output / "test01_move_and_text_demo.pptx")
    return target


def verify_deck(path, data, trees):
    prs = Presentation(path)
    slide = prs.slides[0]
    by_name = {s.name:s for s in slide.shapes}
    for label in data["labels"]:
        assert by_name["Text_" + label["id"]].text == label["text"], label["id"]
        expected = sum(bool(r[1]) for r in label.get("runs", []))
        assert len(by_name["Text_" + label["id"]]._element.xpath(".//a:rPr[@baseline]")) == expected
    for tree in trees:
        node = by_name[tree["id"]]._element
        assert len(node.xpath(".//a:path")) == 1
        assert len(node.xpath(".//a:moveTo")) == len(node.xpath(".//a:close")) == 1
    assert not slide._element.xpath(".//p:pic")
    ids = [n.get("id") for n in slide._element.xpath(".//p:cNvPr")]
    assert len(ids) == len(set(ids)), "Duplicate Office shape IDs"
    return {"verified_text_blocks":len(data["labels"]), "single_shape_trees":len(trees),
            "top_level_objects":len(slide.shapes), "embedded_pictures":0,
            "verification_scope":"Native structure and exact annotated text; visual QA is separate."}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--image", type=Path, required=True)
    parser.add_argument("--annotations", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--trace-python", type=Path, required=True)
    parser.add_argument("--vtracer-packages", type=Path, required=True)
    args = parser.parse_args()
    source = cv2.imread(str(args.image))
    if source is None:
        raise ValueError(f"Cannot read source image: {args.image}")
    size = (source.shape[1], source.shape[0])
    data = json.loads(args.annotations.read_text(encoding="utf-8"))
    validate_annotations(data, size)
    args.output.mkdir(parents=True, exist_ok=True)
    trees = prepare_layers(source, data, args.output)
    trace_background(args.output/"background.png", args.output/"background.svg", args.trace_python, args.vtracer_packages)
    scene = args.output/"scene.svg"
    assemble_scene(args.output/"background.svg", trees, scene, size, data)
    target = make_deck(scene, data, trees, size, args.output)
    report = verify_deck(target, data, trees)
    report["source"] = str(args.image)
    report["tree_annotations"] = trees
    (args.output/"verification.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps({k:v for k,v in report.items() if k != "tree_annotations"}))


if __name__ == "__main__":
    main()
